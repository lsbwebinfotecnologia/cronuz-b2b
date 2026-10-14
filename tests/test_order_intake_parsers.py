import sys
import os
import io

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))

from app.core.order_parsers import parse_text_to_items, parse_spreadsheet_to_items, parse_pdf_to_items
import openpyxl
from fpdf import FPDF


def test_parse_text():
    sample_text = """
    Olá pessoal, segue nosso pedido de reposição:
    9788591851003 - 5 un
    978-85-96008-37-2 qtd: 10
    2x 9788593891007
    ISBN: 9781107556867 | QTD = 3
    9788591851003 2  (adicionando mais 2 do primeiro)
    """

    items = parse_text_to_items(sample_text)
    items_dict = {i["ean"]: i["quantity"] for i in items}

    # 9788591851003 deve somar 5 + 2 = 7
    assert items_dict.get("9788591851003") == 7
    assert items_dict.get("9788596008372") == 10
    assert items_dict.get("9788593891007") == 2
    assert items_dict.get("9781107556867") == 3

    print("✅ test_parse_text passou com sucesso!")


def test_parse_spreadsheet():
    # Cria uma planilha Excel em memória
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Código de Barras", "Título da Obra", "Quantidade Solicitada"])
    ws.append(["9788591851003", "Livro Teste A", 15])
    ws.append(["978-85-96008-37-2", "Livro Teste B", "8"])
    
    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    items = parse_spreadsheet_to_items(excel_bytes, "pedido_livraria.xlsx")
    items_dict = {i["ean"]: i["quantity"] for i in items}

    assert items_dict.get("9788591851003") == 15
    assert items_dict.get("9788596008372") == 8

    print("✅ test_parse_spreadsheet passou com sucesso!")


def test_parse_pdf():
    # Cria um PDF sintético em memória com fpdf2
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text="PEDIDO DE COMPRA - LIVRARIA EXEMPLO", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(200, 10, text="Item 1: ISBN 9788591851003 - Quantidade: 12", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(200, 10, text="Item 2: 978-85-96008-37-2 - Qtd: 6 un", new_x="LMARGIN", new_y="NEXT")
    
    pdf_bytes = bytes(pdf.output())

    items = parse_pdf_to_items(pdf_bytes)
    items_dict = {i["ean"]: i["quantity"] for i in items}

    assert items_dict.get("9788591851003") == 12
    assert items_dict.get("9788596008372") == 6

    print("✅ test_parse_pdf passou com sucesso!")


if __name__ == "__main__":
    test_parse_text()
    test_parse_spreadsheet()
    test_parse_pdf()
    print("\n🎉 TODOS OS TESTES DE PARSERS DE PEDIDOS PASSARAM COM SUCESSO!")
