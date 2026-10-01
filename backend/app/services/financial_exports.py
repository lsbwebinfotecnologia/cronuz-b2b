# -*- coding: utf-8 -*-
"""
Serviço de Exportação de Relatórios Financeiros de Fluxo de Caixa Previsto em Excel (.xlsx) e PDF (.pdf).
Suporta agrupamento por cliente/fornecedor com detalhes minuciosos das observações e composição de lançamentos agrupados.
"""

from io import BytesIO
from datetime import datetime
from typing import List, Optional, Any, Dict
import unicodedata

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from fpdf import FPDF


def _sanitize_text(text: Any) -> str:
    """Higieniza o texto para codificação compatível com o FPDF standard."""
    if text is None:
        return ""
    s = str(text)
    s = unicodedata.normalize('NFKD', s)
    return s.encode('latin-1', 'replace').decode('latin-1')


class FinancialPDFReport(FPDF):
    """Classe base para relatórios em PDF de Fluxo de Caixa."""

    def __init__(self, title: str, subtitle: str, company_name: str, orientation: str = "L"):
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.report_title = title
        self.report_subtitle = subtitle
        self.company_name = company_name

    def header(self):
        page_width = 297 if self.cur_orientation == "L" else 210
        self.set_fill_color(15, 23, 42)
        self.rect(0, 0, page_width, 18, 'F')
        
        self.set_xy(10, 4)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(255, 255, 255)
        comp = _sanitize_text(f"CRONUZ B2B | {self.company_name}")
        self.cell(140, 6, comp, 0, 0, 'L')
        
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(203, 213, 225)
        now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        self.cell(page_width - 160, 6, f"Emitido em: {now_str}", 0, 1, 'R')

        self.set_xy(10, 10)
        self.set_font("Helvetica", "B", 8)
        self.set_text_color(52, 211, 153)
        self.cell(page_width - 20, 5, _sanitize_text(self.report_title.upper()), 0, 1, 'L')

        self.ln(7)

    def footer(self):
        page_width = 297 if self.cur_orientation == "L" else 210
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(148, 163, 184)
        page_str = f"Pagina {self.page_no()}"
        self.cell(0, 6, _sanitize_text(f"{self.report_subtitle}  -  {page_str}"), 0, 0, 'C')


def export_cashflow_forecast_excel(
    customer_groups: List[Dict[str, Any]],
    company_name: str,
    period_label: str = ""
) -> BytesIO:
    """Gera planilha Excel (.xlsx) de Fluxo de Caixa Previsto agrupado por cliente/contato."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fluxo Previsto por Cliente"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=14, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_cust_header = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    font_cust_header = Font(name="Calibri", size=11, bold=True, color="065F46")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_data = Font(name="Calibri", size=9, color="0F172A")
    font_subtotal = Font(name="Calibri", size=9, bold=True, color="0F172A")
    fill_subtotal = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_total = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_total = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws["A1"] = f"CRONUZ B2B — FLUXO DE CAIXA PREVISTO POR CLIENTE / FORNECEDOR"
    ws["A1"].font = font_title
    
    sub = f"Empresa: {company_name}"
    if period_label:
        sub += f" | Período Previsto: {period_label}"
    ws["A2"] = sub
    ws["A2"].font = font_subtitle

    headers = [
        "Nº Parcela", "Tipo", "Categoria", "Conta Prevista", "Vencimento",
        "Valor Previsto (R$)", "Status", "Descrição e Composição (Lançamentos Agrupados)"
    ]

    row_idx = 4
    grand_total_entradas = 0.0
    grand_total_saidas = 0.0
    grand_saldo_liquido = 0.0

    for group in customer_groups:
        cust_name = group.get("customer_name") or "Geral / Sem Cliente Definido"
        cust_doc = group.get("customer_document") or ""
        total_entradas = group.get("total_entradas", 0.0)
        total_saidas = group.get("total_saidas", 0.0)
        saldo_cliente = group.get("saldo_cliente", 0.0)

        grand_total_entradas += total_entradas
        grand_total_saidas += total_saidas
        grand_saldo_liquido += saldo_cliente

        # Cabeçalho do Cliente
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
        cust_cell = ws.cell(row=row_idx, column=1)
        doc_part = f" — CNPJ/CPF: {cust_doc}" if cust_doc else ""
        cust_cell.value = f"CONTATO / CLIENTE: {cust_name}{doc_part}"
        cust_cell.font = font_cust_header
        cust_cell.fill = fill_cust_header
        cust_cell.alignment = Alignment(vertical="center", indent=1)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = border_cell
        row_idx += 1

        # Cabeçalhos de coluna
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=h)
            c.font = font_header
            c.fill = fill_header
            c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 7] else ("right" if col_idx == 6 else "left"), vertical="center")
            c.border = border_cell
        row_idx += 1

        # Linhas de lançamentos
        items = group.get("items", [])
        for i_idx, item in enumerate(items):
            is_even = i_idx % 2 == 0
            row_fill = None if is_even else fill_zebra

            inst_num = f"#{item.get('id')}"
            tipo_desc = "ENTRADA" if item.get("type") == "RECEIVABLE" else "SAÍDA"
            cat_name = item.get("category_name") or "Sem Categoria"
            acc_name = item.get("account_name") or "Padrão"
            due_dt = item.get("due_date") or ""
            val = float(item.get("amount") or 0.0)
            status_desc = item.get("status") or "PENDENTE"
            obs = item.get("description") or ""

            cells_data = [
                (inst_num, "center"),
                (tipo_desc, "center"),
                (cat_name, "left"),
                (acc_name, "left"),
                (due_dt, "center"),
                (val, "right"),
                (status_desc, "center"),
                (obs, "left")
            ]

            for col_idx, (val_data, align) in enumerate(cells_data, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val_data)
                c.font = font_data
                if row_fill:
                    c.fill = row_fill
                c.border = border_cell
                c.alignment = Alignment(horizontal=align, vertical="center")
                if col_idx == 6:
                    c.number_format = '"R$ "#,##0.00'
                    if item.get("type") == "RECEIVABLE":
                        c.font = Font(name="Calibri", size=9, color="059669", bold=True)
                    else:
                        c.font = Font(name="Calibri", size=9, color="DC2626", bold=True)

            row_idx += 1

        # Linha de Subtotal do Cliente
        ws.cell(row=row_idx, column=1, value="SUBTOTAL CONTATO:").font = font_subtotal
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)

        c_val = ws.cell(row=row_idx, column=6, value=saldo_cliente)
        c_val.font = font_subtotal
        c_val.number_format = '"R$ "#,##0.00'
        c_val.alignment = Alignment(horizontal="right")

        ws.cell(row=row_idx, column=7, value=f"Entradas: R$ {total_entradas:,.2f} | Saídas: R$ {total_saidas:,.2f}").font = font_subtitle
        ws.merge_cells(start_row=row_idx, start_column=7, end_row=row_idx, end_column=len(headers))

        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = border_cell
            ws.cell(row=row_idx, column=col).fill = fill_subtotal

        row_idx += 2

    # Linha de Total Geral da Projeção
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    c_tot_label = ws.cell(row=row_idx, column=1, value="SALDO LÍQUIDO PREVISTO TOTAL:")
    c_tot_label.font = font_total
    c_tot_label.alignment = Alignment(horizontal="right", vertical="center")
    
    c_tot_val = ws.cell(row=row_idx, column=6, value=grand_saldo_liquido)
    c_tot_val.font = font_total
    c_tot_val.number_format = '"R$ "#,##0.00'
    c_tot_val.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(start_row=row_idx, start_column=7, end_row=row_idx, end_column=len(headers))
    c_tot_extra = ws.cell(row=row_idx, column=7, value=f"Total Entradas Previstas: R$ {grand_total_entradas:,.2f} | Total Saídas Previstas: R$ {grand_total_saidas:,.2f}")
    c_tot_extra.font = font_total
    c_tot_extra.alignment = Alignment(horizontal="left", vertical="center")

    for col in range(1, len(headers) + 1):
        c = ws.cell(row=row_idx, column=col)
        c.fill = fill_total
        c.border = border_cell

    # Larguras
    col_widths = {1: 12, 2: 12, 3: 24, 4: 20, 5: 14, 6: 18, 7: 14, 8: 55}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_cashflow_forecast_pdf(
    customer_groups: List[Dict[str, Any]],
    company_name: str,
    period_label: str = ""
) -> BytesIO:
    """Gera relatório em PDF de Fluxo de Caixa Previsto com layout financeiro executivo."""
    pdf = FinancialPDFReport(
        title="Projeção de Fluxo de Caixa por Cliente / Fornecedor",
        subtitle=f"Período: {period_label}" if period_label else "Todos os Lançamentos Previstos",
        company_name=company_name,
        orientation="L"
    )
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    grand_total_entradas = 0.0
    grand_total_saidas = 0.0
    grand_saldo_liquido = 0.0

    col_widths = [16, 18, 38, 30, 22, 26, 22, 95]

    for group in customer_groups:
        cust_name = group.get("customer_name") or "Geral / Sem Cliente Definido"
        cust_doc = group.get("customer_document") or ""
        total_entradas = group.get("total_entradas", 0.0)
        total_saidas = group.get("total_saidas", 0.0)
        saldo_cliente = group.get("saldo_cliente", 0.0)

        grand_total_entradas += total_entradas
        grand_total_saidas += total_saidas
        grand_saldo_liquido += saldo_cliente

        # Cabeçalho do Cliente
        pdf.set_fill_color(236, 253, 245)
        pdf.set_text_color(6, 95, 70)
        pdf.set_font("Helvetica", "B", 9)
        doc_part = f" (CNPJ/CPF: {cust_doc})" if cust_doc else ""
        pdf.cell(sum(col_widths), 6.5, _sanitize_text(f" CONTATO: {cust_name}{doc_part}"), border=1, ln=1, fill=True)

        # Cabeçalho de colunas
        pdf.set_fill_color(30, 41, 59)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 7.5)
        headers = ["Nº Parc.", "Tipo", "Categoria", "Conta", "Vencimento", "Valor (R$)", "Status", "Descrição e Composição"]
        for idx, (h, w) in enumerate(zip(headers, col_widths)):
            align = "C" if idx in [0, 1, 4, 6] else ("R" if idx == 5 else "L")
            pdf.cell(w, 5.5, _sanitize_text(h), border=1, align=align, fill=True)
        pdf.ln()

        # Linhas de lançamentos
        items = group.get("items", [])
        pdf.set_font("Helvetica", "", 7.5)
        for i_idx, item in enumerate(items):
            if i_idx % 2 == 1:
                pdf.set_fill_color(248, 250, 252)
                fill = True
            else:
                fill = False

            pdf.set_text_color(15, 23, 42)
            inst_num = f"#{item.get('id')}"
            tipo = "RECEITA" if item.get("type") == "RECEIVABLE" else "DESPESA"
            cat_name = item.get("category_name") or "Sem Categoria"
            if len(cat_name) > 22:
                cat_name = cat_name[:20] + ".."
            acc_name = item.get("account_name") or "Padrão"
            if len(acc_name) > 18:
                acc_name = acc_name[:16] + ".."
            due_dt = item.get("due_date") or ""
            val_str = f"{float(item.get('amount') or 0.0):,.2f}"
            status_desc = item.get("status") or "PENDENTE"
            obs = item.get("description") or ""
            if len(obs) > 60:
                obs = obs[:58] + ".."

            row_data = [inst_num, tipo, cat_name, acc_name, due_dt, val_str, status_desc, obs]
            for idx, (val, w) in enumerate(zip(row_data, col_widths)):
                align = "C" if idx in [0, 1, 4, 6] else ("R" if idx == 5 else "L")
                # Cor do valor
                if idx == 5:
                    if item.get("type") == "RECEIVABLE":
                        pdf.set_text_color(5, 150, 105)
                    else:
                        pdf.set_text_color(220, 38, 38)
                else:
                    pdf.set_text_color(15, 23, 42)
                pdf.cell(w, 5, _sanitize_text(val), border=1, align=align, fill=fill)
            pdf.ln()

        # Subtotal do Contato
        pdf.set_fill_color(241, 245, 249)
        pdf.set_font("Helvetica", "B", 7.5)
        pdf.set_text_color(30, 41, 59)
        pdf.cell(sum(col_widths[:5]), 5.5, _sanitize_text(f"Subtotal: Entradas R$ {total_entradas:,.2f} | Saídas R$ {total_saidas:,.2f}"), border=1, align="R", fill=True)
        pdf.cell(col_widths[5], 5.5, f"R$ {saldo_cliente:,.2f}", border=1, align="R", fill=True)
        pdf.cell(sum(col_widths[6:]), 5.5, "", border=1, fill=True)
        pdf.ln(7)

    # Total Geral
    pdf.set_fill_color(15, 23, 42)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.cell(sum(col_widths[:5]), 6.5, _sanitize_text(f"TOTAL PREVISTO (Entradas: R$ {grand_total_entradas:,.2f} | Saídas: R$ {grand_total_saidas:,.2f})"), border=1, align="R", fill=True)
    pdf.cell(col_widths[5], 6.5, f"R$ {grand_saldo_liquido:,.2f}", border=1, align="R", fill=True)
    pdf.cell(sum(col_widths[6:]), 6.5, "", border=1, fill=True)

    buffer = BytesIO()
    out = pdf.output()
    pdf_bytes = bytes(out) if not isinstance(out, (bytes, bytearray)) else bytes(out)
    buffer.write(pdf_bytes)
    buffer.seek(0)
    return buffer
