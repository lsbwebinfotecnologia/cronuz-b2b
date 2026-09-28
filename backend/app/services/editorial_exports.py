# -*- coding: utf-8 -*-
"""
Serviço de Exportação de Relatórios Editoriais em Excel (.xlsx) e PDF (.pdf).
Suporta Relatório de Serviços & Custos e Relatório de Movimentações & Linha do Tempo.
Suporta escopo por projeto específico ou geral da empresa.
"""

from io import BytesIO
from datetime import datetime
from typing import List, Optional, Any, Dict
import unicodedata

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from fpdf import FPDF


def _sanitize_text(text: Any) -> str:
    """Higieniza o texto para codificação compatível com o FPDF standard."""
    if text is None:
        return ""
    s = str(text)
    s = unicodedata.normalize('NFKD', s)
    return s.encode('latin-1', 'replace').decode('latin-1')


class EditorialPDFReport(FPDF):
    """Classe base para relatórios em PDF com cabeçalho e rodapé elegantes."""

    def __init__(self, title: str, subtitle: str, company_name: str, orientation: str = "L"):
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.report_title = title
        self.report_subtitle = subtitle
        self.company_name = company_name

    def header(self):
        page_width = 297 if self.cur_orientation == "L" else 210
        self.set_fill_color(30, 41, 59)
        self.rect(0, 0, page_width, 18, 'F')
        
        self.set_xy(10, 4)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(255, 255, 255)
        comp = _sanitize_text(f"CRONUZ B2B | {self.company_name}")
        self.cell(140, 6, comp, 0, 0, 'L')
        
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(203, 213, 225)
        now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        self.cell(page_width - 160, 6, f"Emitido em: {now_str}", 0, 1, 'R')

        self.set_xy(10, 10)
        self.set_font("Helvetica", "B", 8)
        self.set_text_color(147, 197, 253)
        self.cell(page_width - 20, 5, _sanitize_text(self.report_title.upper()), 0, 1, 'L')

        self.ln(7)

    def footer(self):
        page_width = 297 if self.cur_orientation == "L" else 210
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(148, 163, 184)
        page_str = f"Pagina {self.page_no()}"
        self.cell(0, 6, _sanitize_text(f"{self.report_subtitle}  -  {page_str}"), 0, 0, 'C')


def export_services_excel(
    costs_data: List[Dict[str, Any]],
    company_name: str,
    project_title: Optional[str] = None
) -> BytesIO:
    """Gera planilha Excel (.xlsx) altamente formatada com grade de serviços e custos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Servicos e Custos"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_total = Font(name="Calibri", size=10, bold=True, color="0F172A")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws["A1"] = f"CRONUZ B2B — RELATORIO DE SERVICOS E CUSTOS EDITORIAIS"
    ws["A1"].font = font_title
    
    sub = f"Empresa: {company_name} | "
    if project_title:
        sub += f"Filtro: Projeto Especifico ({project_title})"
    else:
        sub += "Filtro: Geral (Todos os Projetos)"
    sub += f" | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"] = sub
    ws["A2"].font = font_subtitle

    headers = [
        "ID",
        "Projeto / Obra",
        "Tipo de Servico",
        "Descricao Detalhada",
        "Prestador de Servico",
        "Contato / Chave Pix",
        "Etapa Vinculada",
        "Unidade",
        "Qtd",
        "Valor Unit. (R$)",
        "Valor Orcado (R$)",
        "Valor Real (R$)",
        "Status Pagamento",
        "Data Cadastro",
        "Data Pagamento",
        "No NF-e / Recibo"
    ]

    start_row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell

    ws.row_dimensions[start_row].height = 28

    current_row = start_row + 1
    total_orcado_sum = 0.0
    total_real_sum = 0.0

    for item in costs_data:
        is_zebra = (current_row % 2 == 0)
        row_fill = fill_zebra if is_zebra else None

        row_values = [
            item.get("id"),
            item.get("project_title", "N/D"),
            item.get("service_type", "").replace("_", " "),
            item.get("description", ""),
            item.get("professional_name", "Nao alocado"),
            item.get("professional_pix", "-"),
            item.get("stage_name", "Geral"),
            item.get("unit_type", ""),
            item.get("quantity", 1),
            item.get("unit_value", 0.0),
            item.get("estimated_total", 0.0),
            item.get("actual_total", 0.0),
            item.get("payment_status", "ORCADO"),
            item.get("created_at_str", "-"),
            item.get("paid_at_str", "-"),
            item.get("invoice_number", "-")
        ]

        total_orcado_sum += float(item.get("estimated_total", 0.0) or 0.0)
        total_real_sum += float(item.get("actual_total", 0.0) or 0.0)

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if row_fill:
                cell.fill = row_fill

            if col_idx in [1, 8, 9, 13, 14, 15]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [10, 11, 12]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '"R$" #,##0.00'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Linha de Totais
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    total_label = ws.cell(row=current_row, column=1, value="TOTAIS CONSOLIDADOS")
    total_label.font = font_total
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    total_label.fill = fill_total
    total_label.border = border_cell

    for c in range(1, 11):
        ws.cell(row=current_row, column=c).border = border_cell
        ws.cell(row=current_row, column=c).fill = fill_total

    cell_orc = ws.cell(row=current_row, column=11, value=total_orcado_sum)
    cell_orc.font = font_total
    cell_orc.alignment = Alignment(horizontal="right", vertical="center")
    cell_orc.number_format = '"R$" #,##0.00'
    cell_orc.fill = fill_total
    cell_orc.border = border_cell

    cell_real = ws.cell(row=current_row, column=12, value=total_real_sum)
    cell_real.font = font_total
    cell_real.alignment = Alignment(horizontal="right", vertical="center")
    cell_real.number_format = '"R$" #,##0.00'
    cell_real.fill = fill_total
    cell_real.border = border_cell

    for c in range(13, 17):
        ws.cell(row=current_row, column=c).border = border_cell
        ws.cell(row=current_row, column=c).fill = fill_total

    ws.row_dimensions[current_row].height = 24

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.row in [1, 2]:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_services_pdf(
    costs_data: List[Dict[str, Any]],
    company_name: str,
    project_title: Optional[str] = None
) -> bytes:
    """Gera relatório em PDF elegante (Paisagem A4) com os serviços e custos."""
    title = "Relatorio de Servicos & Custos Editoriais"
    scope_str = f"Projeto: {project_title}" if project_title else "Todos os Projetos (Geral)"
    
    pdf = EditorialPDFReport(
        title=title,
        subtitle=f"{company_name} | {scope_str}",
        company_name=company_name,
        orientation="L"
    )
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(30, 41, 59)
    pdf.cell(0, 6, _sanitize_text(title), 0, 1, 'L')
    
    pdf.set_font("Helvetica", "", 8.5)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, _sanitize_text(f"Filtro: {scope_str}  |  Registros listados: {len(costs_data)}"), 0, 1, 'L')
    pdf.ln(3)

    cols = [
        {"name": "Obra / Projeto", "w": 45, "align": "L"},
        {"name": "Servico", "w": 32, "align": "L"},
        {"name": "Descricao / Detalhes", "w": 55, "align": "L"},
        {"name": "Prestador / Pix", "w": 42, "align": "L"},
        {"name": "Qtd/Un", "w": 20, "align": "C"},
        {"name": "Orcado", "w": 25, "align": "R"},
        {"name": "Realizado", "w": 25, "align": "R"},
        {"name": "Status", "w": 33, "align": "C"}
    ]

    pdf.set_fill_color(51, 65, 85)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    for c in cols:
        pdf.cell(c["w"], 7, _sanitize_text(c["name"]), 1, 0, c["align"], fill=True)
    pdf.ln(7)

    total_orc = 0.0
    total_real = 0.0

    pdf.set_font("Helvetica", "", 7.5)
    for idx, item in enumerate(costs_data):
        is_even = (idx % 2 == 0)
        if is_even:
            pdf.set_fill_color(248, 250, 252)
        else:
            pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(30, 41, 59)

        orc = float(item.get("estimated_total", 0.0) or 0.0)
        real = float(item.get("actual_total", 0.0) or 0.0)
        total_orc += orc
        total_real += real

        proj = item.get("project_title", "N/D")[:25]
        serv = item.get("service_type", "").replace("_", " ")[:18]
        desc = item.get("description", "")[:35]
        prof = item.get("professional_name") or "Nao alocado"
        if item.get("professional_pix"):
            prof = f"{prof} ({item.get('professional_pix')[:14]})"
        prof = prof[:26]

        qtd_un = f"{item.get('quantity', 1)} {item.get('unit_type', '')}"[:12]
        status = item.get("payment_status", "ORCADO")

        pdf.cell(cols[0]["w"], 6, _sanitize_text(proj), 1, 0, 'L', fill=True)
        pdf.cell(cols[1]["w"], 6, _sanitize_text(serv), 1, 0, 'L', fill=True)
        pdf.cell(cols[2]["w"], 6, _sanitize_text(desc), 1, 0, 'L', fill=True)
        pdf.cell(cols[3]["w"], 6, _sanitize_text(prof), 1, 0, 'L', fill=True)
        pdf.cell(cols[4]["w"], 6, _sanitize_text(qtd_un), 1, 0, 'C', fill=True)
        pdf.cell(cols[5]["w"], 6, f"R$ {orc:,.2f}", 1, 0, 'R', fill=True)
        pdf.cell(cols[6]["w"], 6, f"R$ {real:,.2f}", 1, 0, 'R', fill=True)
        pdf.cell(cols[7]["w"], 6, _sanitize_text(status), 1, 1, 'C', fill=True)

    pdf.set_fill_color(226, 232, 240)
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_text_color(15, 23, 42)
    tot_label_w = sum(c["w"] for c in cols[:5])
    pdf.cell(tot_label_w, 7, "TOTAIS:", 1, 0, 'R', fill=True)
    pdf.cell(cols[5]["w"], 7, f"R$ {total_orc:,.2f}", 1, 0, 'R', fill=True)
    pdf.cell(cols[6]["w"], 7, f"R$ {total_real:,.2f}", 1, 0, 'R', fill=True)
    pdf.cell(cols[7]["w"], 7, "", 1, 1, 'C', fill=True)

    return bytes(pdf.output())


def export_movements_excel(
    movements_data: List[Dict[str, Any]],
    company_name: str,
    project_title: Optional[str] = None
) -> BytesIO:
    """Gera planilha Excel (.xlsx) com o histórico e movimentações das etapas dos projetos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimentacoes de Projetos"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_data = Font(name="Calibri", size=10, color="0F172A")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws["A1"] = f"CRONUZ B2B — RELATORIO DE MOVIMENTACOES & HISTORICO EDITORIAL"
    ws["A1"].font = font_title
    
    sub = f"Empresa: {company_name} | "
    if project_title:
        sub += f"Filtro: Projeto Especifico ({project_title})"
    else:
        sub += "Filtro: Geral (Todos os Projetos)"
    sub += f" | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"] = sub
    ws["A2"].font = font_subtitle

    headers = [
        "Data & Hora",
        "Projeto / Obra",
        "Codigo Interno",
        "Etapa de Origem",
        "Etapa de Destino",
        "Acao Executada",
        "Usuario Responsavel",
        "Anotacoes & Justificativas"
    ]

    start_row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell

    ws.row_dimensions[start_row].height = 28

    current_row = start_row + 1
    for item in movements_data:
        is_zebra = (current_row % 2 == 0)
        row_fill = fill_zebra if is_zebra else None

        row_values = [
            item.get("created_at_str", "-"),
            item.get("project_title", "N/D"),
            f"#{item.get('project_local_id', item.get('project_id', '-'))}",
            item.get("from_stage_name", "Inicio do Fluxo"),
            item.get("to_stage_name", "Etapa Inicial"),
            item.get("action", "MUDANCA_ETAPA"),
            item.get("user_name", "Sistema"),
            item.get("notes", "-")
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if row_fill:
                cell.fill = row_fill

            if col_idx in [1, 3, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.row in [1, 2]:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_movements_pdf(
    movements_data: List[Dict[str, Any]],
    company_name: str,
    project_title: Optional[str] = None
) -> bytes:
    """Gera relatório em PDF elegante (Paisagem A4) com a linha do tempo e movimentações."""
    title = "Relatorio de Movimentacoes & Timeline Editorial"
    scope_str = f"Projeto: {project_title}" if project_title else "Todos os Projetos (Geral)"
    
    pdf = EditorialPDFReport(
        title=title,
        subtitle=f"{company_name} | {scope_str}",
        company_name=company_name,
        orientation="L"
    )
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(30, 41, 59)
    pdf.cell(0, 6, _sanitize_text(title), 0, 1, 'L')
    
    pdf.set_font("Helvetica", "", 8.5)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, _sanitize_text(f"Filtro: {scope_str}  |  Movimentacoes registradas: {len(movements_data)}"), 0, 1, 'L')
    pdf.ln(3)

    cols = [
        {"name": "Data & Hora", "w": 30, "align": "C"},
        {"name": "Projeto / Obra", "w": 52, "align": "L"},
        {"name": "De Etapa", "w": 40, "align": "L"},
        {"name": "Para Etapa", "w": 40, "align": "L"},
        {"name": "Acao", "w": 30, "align": "C"},
        {"name": "Usuario", "w": 35, "align": "L"},
        {"name": "Observacoes / Notas", "w": 50, "align": "L"}
    ]

    pdf.set_fill_color(30, 58, 138)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    for c in cols:
        pdf.cell(c["w"], 7, _sanitize_text(c["name"]), 1, 0, c["align"], fill=True)
    pdf.ln(7)

    pdf.set_font("Helvetica", "", 7.5)
    for idx, item in enumerate(movements_data):
        is_even = (idx % 2 == 0)
        if is_even:
            pdf.set_fill_color(248, 250, 252)
        else:
            pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(30, 41, 59)

        d_str = item.get("created_at_str", "-")[:16]
        proj = item.get("project_title", "N/D")[:30]
        from_st = item.get("from_stage_name", "Inicio")[:22]
        to_st = item.get("to_stage_name", "Etapa")[:22]
        action = item.get("action", "MUDANCA")[:16]
        user_name = item.get("user_name", "Sistema")[:20]
        notes = item.get("notes", "-")[:32]

        pdf.cell(cols[0]["w"], 6, _sanitize_text(d_str), 1, 0, 'C', fill=True)
        pdf.cell(cols[1]["w"], 6, _sanitize_text(proj), 1, 0, 'L', fill=True)
        pdf.cell(cols[2]["w"], 6, _sanitize_text(from_st), 1, 0, 'L', fill=True)
        pdf.cell(cols[3]["w"], 6, _sanitize_text(to_st), 1, 0, 'L', fill=True)
        pdf.cell(cols[4]["w"], 6, _sanitize_text(action), 1, 0, 'C', fill=True)
        pdf.cell(cols[5]["w"], 6, _sanitize_text(user_name), 1, 0, 'L', fill=True)
        pdf.cell(cols[6]["w"], 6, _sanitize_text(notes), 1, 1, 'L', fill=True)

    return bytes(pdf.output())
