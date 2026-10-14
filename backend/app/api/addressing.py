import csv
import io
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
import openpyxl

from app.core.dependencies import get_current_user
from app.core.upload_security import validate_file_size_and_extension
from app.db.session import get_db
from app.models.company import Company
from app.models.product_address import ProductAddress
from app.models.seller_branch import SellerBranch
from app.models.user import User
from app.schemas.product_address import (
    ProductAddressCreate,
    ProductAddressResponse,
    ProductAddressSyncBatchRequest,
    ProductAddressSyncBatchResponse,
    ProductAddressUpdate,
    ProductAddressUploadSummary,
)

router = APIRouter(prefix="/addressing", tags=["addressing"])


def _get_user_company_id(current_user: User) -> int:
    if not current_user.company_id and current_user.type == "MASTER":
        return 1
    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Usuário sem empresa vinculada.",
        )
    return current_user.company_id


@router.get("", response_model=dict)
def list_addresses(
    branch_id: Optional[int] = Query(None),
    cod_local_estoque: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Lista os endereçamentos cadastrados com filtros e paginação."""
    company_id = _get_user_company_id(current_user)

    query = (
        db.query(
            ProductAddress,
            SellerBranch.nome.label("branch_name"),
            User.name.label("creator_name"),
        )
        .outerjoin(SellerBranch, ProductAddress.branch_id == SellerBranch.id)
        .outerjoin(User, ProductAddress.created_by_user_id == User.id)
        .filter(ProductAddress.company_id == company_id)
    )

    if branch_id:
        query = query.filter(ProductAddress.branch_id == branch_id)

    if cod_local_estoque:
        query = query.filter(ProductAddress.cod_local_estoque == cod_local_estoque)

    if search:
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                ProductAddress.isbn.ilike(search_term),
                ProductAddress.horus_cod_item.ilike(search_term),
                ProductAddress.title.ilike(search_term),
                ProductAddress.address.ilike(search_term),
                ProductAddress.notes.ilike(search_term),
            )
        )

    total = query.count()
    items_raw = (
        query.order_by(ProductAddress.updated_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )

    items = []
    for addr, branch_name, creator_name in items_raw:
        items.append(
            {
                "id": addr.id,
                "company_id": addr.company_id,
                "branch_id": addr.branch_id,
                "cod_local_estoque": addr.cod_local_estoque,
                "isbn": addr.isbn,
                "horus_cod_item": addr.horus_cod_item,
                "title": addr.title,
                "address": addr.address,
                "notes": addr.notes,
                "created_by_user_id": addr.created_by_user_id,
                "created_at": addr.created_at,
                "updated_at": addr.updated_at,
                "branch_name": branch_name,
                "creator_name": creator_name,
            }
        )

    return {
        "items": items,
        "total": total,
        "page": page,
        "per_page": per_page,
        "pages": (total + per_page - 1) // per_page if per_page else 1,
    }


@router.get("/by-product")
def get_addresses_by_product(
    isbn: Optional[str] = Query(None),
    horus_cod_item: Optional[str] = Query(None),
    branch_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Retorna os endereços de um produto específico (pode ter múltiplos endereços)."""
    company_id = _get_user_company_id(current_user)

    if not isbn and not horus_cod_item:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Informe o ISBN ou o Código do Item para consulta de endereços.",
        )

    filters = [ProductAddress.company_id == company_id]
    if branch_id:
        filters.append(ProductAddress.branch_id == branch_id)

    match_conditions = []
    if isbn:
        clean_isbn = isbn.strip().replace("-", "")
        match_conditions.append(ProductAddress.isbn == clean_isbn)
        match_conditions.append(ProductAddress.isbn == isbn.strip())
    if horus_cod_item:
        match_conditions.append(ProductAddress.horus_cod_item == str(horus_cod_item).strip())

    filters.append(or_(*match_conditions))

    results = (
        db.query(
            ProductAddress,
            SellerBranch.nome.label("branch_name"),
            SellerBranch.cod_empresa.label("branch_cod_empresa"),
            SellerBranch.cod_filial.label("branch_cod_filial"),
        )
        .outerjoin(SellerBranch, ProductAddress.branch_id == SellerBranch.id)
        .filter(*filters)
        .order_by(ProductAddress.address.asc())
        .all()
    )

    addresses = []
    for addr, b_name, b_cod_empresa, b_cod_filial in results:
        addresses.append(
            {
                "id": addr.id,
                "branch_id": addr.branch_id,
                "branch_name": b_name,
                "branch_cod_empresa": b_cod_empresa,
                "branch_cod_filial": b_cod_filial,
                "cod_local_estoque": addr.cod_local_estoque,
                "isbn": addr.isbn,
                "horus_cod_item": addr.horus_cod_item,
                "title": addr.title,
                "address": addr.address,
                "notes": addr.notes,
                "updated_at": addr.updated_at,
            }
        )

    return {"addresses": addresses, "total": len(addresses)}


@router.post("", response_model=dict)
def create_address(
    data: ProductAddressCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Cadastra ou atualiza um endereço de produto individual."""
    company_id = _get_user_company_id(current_user)

    # Verifica se a filial existe e pertence à empresa
    branch = (
        db.query(SellerBranch)
        .filter(SellerBranch.id == data.branch_id, SellerBranch.company_id == company_id)
        .first()
    )
    if not branch:
        raise HTTPException(status_code=404, detail="Filial não encontrada.")

    clean_isbn = data.isbn.strip().replace("-", "")
    clean_address = data.address.strip().upper()
    clean_local = data.cod_local_estoque.strip() if data.cod_local_estoque else None

    # Busca se já existe exatamente o mesmo endereço
    existing = (
        db.query(ProductAddress)
        .filter(
            ProductAddress.company_id == company_id,
            ProductAddress.branch_id == data.branch_id,
            func.coalesce(ProductAddress.cod_local_estoque, "") == (clean_local or ""),
            ProductAddress.isbn == clean_isbn,
            ProductAddress.address == clean_address,
        )
        .first()
    )

    now = datetime.utcnow()
    if existing:
        if data.notes is not None:
            existing.notes = data.notes
        if data.title:
            existing.title = data.title
        if data.horus_cod_item:
            existing.horus_cod_item = data.horus_cod_item
        existing.updated_at = now
        db.commit()
        db.refresh(existing)
        return {"message": "Endereço já existia e foi atualizado.", "id": existing.id, "created": False}

    new_addr = ProductAddress(
        company_id=company_id,
        branch_id=data.branch_id,
        cod_local_estoque=clean_local,
        isbn=clean_isbn,
        horus_cod_item=data.horus_cod_item.strip() if data.horus_cod_item else None,
        title=data.title.strip() if data.title else None,
        address=clean_address,
        notes=data.notes.strip() if data.notes else None,
        created_by_user_id=current_user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(new_addr)
    db.commit()
    db.refresh(new_addr)
    return {"message": "Endereço cadastrado com sucesso.", "id": new_addr.id, "created": True}


@router.put("/{address_id}", response_model=dict)
def update_address(
    address_id: int,
    data: ProductAddressUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Atualiza dados de um endereço existente."""
    company_id = _get_user_company_id(current_user)

    addr = (
        db.query(ProductAddress)
        .filter(ProductAddress.id == address_id, ProductAddress.company_id == company_id)
        .first()
    )
    if not addr:
        raise HTTPException(status_code=404, detail="Registro de endereço não encontrado.")

    if data.branch_id is not None:
        branch = (
            db.query(SellerBranch)
            .filter(SellerBranch.id == data.branch_id, SellerBranch.company_id == company_id)
            .first()
        )
        if not branch:
            raise HTTPException(status_code=404, detail="Filial não encontrada.")
        addr.branch_id = data.branch_id

    if data.address is not None:
        addr.address = data.address.strip().upper()
    if data.cod_local_estoque is not None:
        addr.cod_local_estoque = data.cod_local_estoque.strip() if data.cod_local_estoque else None
    if data.isbn is not None:
        addr.isbn = data.isbn.strip().replace("-", "")
    if data.horus_cod_item is not None:
        addr.horus_cod_item = data.horus_cod_item.strip() if data.horus_cod_item else None
    if data.title is not None:
        addr.title = data.title.strip() if data.title else None
    if data.notes is not None:
        addr.notes = data.notes.strip() if data.notes else None

    addr.updated_at = datetime.utcnow()
    db.commit()
    return {"message": "Endereço atualizado com sucesso.", "id": addr.id}


@router.delete("/{address_id}")
def delete_address(
    address_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Exclui um registro de endereço."""
    company_id = _get_user_company_id(current_user)

    addr = (
        db.query(ProductAddress)
        .filter(ProductAddress.id == address_id, ProductAddress.company_id == company_id)
        .first()
    )
    if not addr:
        raise HTTPException(status_code=404, detail="Registro de endereço não encontrado.")

    db.delete(addr)
    db.commit()
    return {"message": "Endereço removido com sucesso."}


@router.delete("/batch/delete")
def batch_delete_addresses(
    ids: List[int],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Exclui múltiplos endereços em lote."""
    company_id = _get_user_company_id(current_user)

    deleted = (
        db.query(ProductAddress)
        .filter(ProductAddress.company_id == company_id, ProductAddress.id.in_(ids))
        .delete(synchronize_session=False)
    )
    db.commit()
    return {"message": f"{deleted} endereços removidos com sucesso.", "deleted_count": deleted}


@router.post("/sync-batch", response_model=ProductAddressSyncBatchResponse)
def sync_batch_addresses(
    payload: ProductAddressSyncBatchRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Sincroniza em lote os itens bipados no modo mobile/offline.
    Pode adicionar (append) ou atualizar registros existentes.
    """
    company_id = _get_user_company_id(current_user)

    if not payload.items:
        return ProductAddressSyncBatchResponse(
            inserted_count=0, updated_count=0, total_received=0, errors=[]
        )

    inserted = 0
    updated = 0
    errors = []
    now = datetime.utcnow()

    # Cache de validação de filiais da empresa
    valid_branch_ids = set(
        b.id for b in db.query(SellerBranch.id).filter(SellerBranch.company_id == company_id).all()
    )

    for idx, item in enumerate(payload.items):
        try:
            if item.branch_id not in valid_branch_ids:
                errors.append(f"Linha {idx + 1}: Filial {item.branch_id} inválida ou não pertence à empresa.")
                continue

            clean_isbn = item.isbn.strip().replace("-", "")
            clean_address = item.address.strip().upper()
            clean_local = item.cod_local_estoque.strip() if item.cod_local_estoque else None

            if not clean_isbn or not clean_address:
                errors.append(f"Linha {idx + 1}: ISBN e Endereço são obrigatórios.")
                continue

            existing = (
                db.query(ProductAddress)
                .filter(
                    ProductAddress.company_id == company_id,
                    ProductAddress.branch_id == item.branch_id,
                    func.coalesce(ProductAddress.cod_local_estoque, "") == (clean_local or ""),
                    ProductAddress.isbn == clean_isbn,
                    ProductAddress.address == clean_address,
                )
                .first()
            )

            if existing:
                if item.notes:
                    existing.notes = item.notes
                if item.title and not existing.title:
                    existing.title = item.title
                if item.horus_cod_item and not existing.horus_cod_item:
                    existing.horus_cod_item = item.horus_cod_item
                existing.updated_at = now
                updated += 1
            else:
                new_record = ProductAddress(
                    company_id=company_id,
                    branch_id=item.branch_id,
                    cod_local_estoque=clean_local,
                    isbn=clean_isbn,
                    horus_cod_item=item.horus_cod_item.strip() if item.horus_cod_item else None,
                    title=item.title.strip() if item.title else None,
                    address=clean_address,
                    notes=item.notes.strip() if item.notes else None,
                    created_by_user_id=current_user.id,
                    created_at=now,
                    updated_at=now,
                )
                db.add(new_record)
                inserted += 1

        except Exception as e:
            errors.append(f"Linha {idx + 1} ({getattr(item, 'isbn', 'item')}): {str(e)}")

    db.commit()

    return ProductAddressSyncBatchResponse(
        inserted_count=inserted,
        updated_count=updated,
        total_received=len(payload.items),
        errors=errors[:50],  # Limita mensagens de erro para segurança de payload
    )


@router.get("/template")
def download_addressing_template():
    """Gera e faz download de uma planilha modelo (.xlsx) para importação de endereços."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enderecamento"

    headers = [
        "ISBN",
        "ENDERECO",
        "TITULO",
        "COD_ITEM",
        "COD_LOCAL_ESTOQUE",
        "OBSERVACOES",
    ]
    ws.append(headers)

    # Exemplos
    ws.append(["9786555640000", "RUA-A-P01", "Livro Exemplo 1", "102030", "001", "Frente de loja"])
    ws.append(["9786555640000", "ESTOQUE-B-05", "Livro Exemplo 1", "102030", "001", "Estoque reserva"])
    ws.append(["9788535914849", "CORREDOR-3-C", "Livro Exemplo 2", "104050", "001", ""])

    # Ajusta largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = "modelo_enderecamento_produtos.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/upload", response_model=ProductAddressUploadSummary)
async def upload_addressing_sheet(
    file: UploadFile = File(...),
    branch_id: int = Form(...),
    cod_local_estoque: Optional[str] = Form(None),
    mode: str = Form("append"),  # "append" ou "replace"
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Importa planilha Excel (.xlsx) ou CSV para endereçamento de produtos.
    Validação mandatória de até 10 MB.
    Suporta append ou overwrite (substituição).
    """
    company_id = _get_user_company_id(current_user)

    # 1. Valida tamanho e extensão (categoria 'sheet' limita a 10MB)
    validate_file_size_and_extension(file, category="sheet")

    # 2. Valida se a filial informada pertence à empresa
    branch = (
        db.query(SellerBranch)
        .filter(SellerBranch.id == branch_id, SellerBranch.company_id == company_id)
        .first()
    )
    if not branch:
        raise HTTPException(status_code=404, detail="Filial selecionada não existe ou não pertence à sua empresa.")

    content = await file.read()
    filename_lower = (file.filename or "").lower()

    rows = []
    # 3. Processamento de Excel ou CSV
    if filename_lower.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        sheet = wb.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        if not raw_rows:
            raise HTTPException(status_code=400, detail="A planilha está vazia.")

        headers = [str(c or "").strip().upper() for c in raw_rows[0]]
        for r in raw_rows[1:]:
            if not any(r):
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                if idx < len(r):
                    val = r[idx]
                    row_dict[h] = str(val).strip() if val is not None else ""
            rows.append(row_dict)
    else:
        # CSV
        text_data = content.decode("utf-8-sig", errors="replace")
        sample = text_data[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ";" if ";" in sample else ","

        reader = csv.DictReader(io.StringIO(text_data), delimiter=delimiter)
        for r in reader:
            clean_r = {str(k or "").strip().upper(): str(v or "").strip() for k, v in r.items() if k}
            if any(clean_r.values()):
                rows.append(clean_r)

    if not rows:
        raise HTTPException(status_code=400, detail="Nenhum registro encontrado no arquivo.")

    # 4. Modo Replace (Substituição)
    clean_default_local = cod_local_estoque.strip() if cod_local_estoque else None
    if mode == "replace":
        del_query = db.query(ProductAddress).filter(
            ProductAddress.company_id == company_id,
            ProductAddress.branch_id == branch_id,
        )
        if clean_default_local:
            del_query = del_query.filter(ProductAddress.cod_local_estoque == clean_default_local)
        del_query.delete(synchronize_session=False)
        db.commit()

    # 5. Processamento dos registros
    total_rows = len(rows)
    inserted = 0
    updated = 0
    skipped = 0
    errors = []
    now = datetime.utcnow()

    def _find_field(row: dict, keys: List[str]) -> Optional[str]:
        for k in keys:
            for actual_k, val in row.items():
                if k in actual_k:
                    return val
        return None

    isbn_keys = ["ISBN", "BARRAS", "CODIGO DE BARRAS", "COD_BARRA", "EAN"]
    addr_keys = ["ENDERECO", "ENDEREÇO", "LOCALIZACAO", "LOCALIZAÇÃO", "POSICAO", "POSIÇÃO"]
    title_keys = ["TITULO", "TÍTULO", "NOME", "DESCRICAO", "DESCRIÇÃO"]
    cod_item_keys = ["COD_ITEM", "CODIGO ITEM", "CÓDIGO ITEM", "COD ITEM"]
    local_keys = ["COD_LOCAL_ESTOQUE", "LOCAL_ESTOQUE", "LOCAL DE ESTOQUE", "LOCAL ESTOQUE"]
    notes_keys = ["OBSERVACOES", "OBSERVAÇÕES", "NOTAS", "OBS"]

    for idx, r in enumerate(rows):
        line_num = idx + 2
        raw_isbn = _find_field(r, isbn_keys)
        raw_addr = _find_field(r, addr_keys)

        if not raw_isbn or not raw_addr:
            skipped += 1
            if len(errors) < 30:
                errors.append(f"Linha {line_num}: ISBN e Endereço são obrigatórios.")
            continue

        clean_isbn = raw_isbn.replace("-", "").strip()
        # Tratamento de número com ponto vindo de excel tipo 9786555640000.0
        if clean_isbn.endswith(".0"):
            clean_isbn = clean_isbn[:-2]

        clean_addr = raw_addr.strip().upper()
        if not clean_isbn or not clean_addr:
            skipped += 1
            continue

        row_local = _find_field(r, local_keys) or clean_default_local
        if row_local:
            row_local = row_local.strip()
            if row_local.endswith(".0"):
                row_local = row_local[:-2]

        raw_title = _find_field(r, title_keys)
        clean_title = raw_title.strip() if raw_title else None

        raw_cod_item = _find_field(r, cod_item_keys)
        clean_cod_item = raw_cod_item.strip() if raw_cod_item else None
        if clean_cod_item and clean_cod_item.endswith(".0"):
            clean_cod_item = clean_cod_item[:-2]

        raw_notes = _find_field(r, notes_keys)
        clean_notes = raw_notes.strip() if raw_notes else None

        existing = (
            db.query(ProductAddress)
            .filter(
                ProductAddress.company_id == company_id,
                ProductAddress.branch_id == branch_id,
                func.coalesce(ProductAddress.cod_local_estoque, "") == (row_local or ""),
                ProductAddress.isbn == clean_isbn,
                ProductAddress.address == clean_addr,
            )
            .first()
        )

        if existing:
            if clean_notes:
                existing.notes = clean_notes
            if clean_title:
                existing.title = clean_title
            if clean_cod_item:
                existing.horus_cod_item = clean_cod_item
            existing.updated_at = now
            updated += 1
        else:
            new_record = ProductAddress(
                company_id=company_id,
                branch_id=branch_id,
                cod_local_estoque=row_local,
                isbn=clean_isbn,
                horus_cod_item=clean_cod_item,
                title=clean_title,
                address=clean_addr,
                notes=clean_notes,
                created_by_user_id=current_user.id,
                created_at=now,
                updated_at=now,
            )
            db.add(new_record)
            inserted += 1

    db.commit()

    return ProductAddressUploadSummary(
        total_rows=total_rows,
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        errors=errors[:30],
    )
