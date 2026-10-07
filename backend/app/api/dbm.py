"""
backend/app/api/dbm.py
----------------------
Endpoints especializados para o módulo DBM (Painel Operacional & CRM):
- Cadastro e consulta de empresas integradas à base Cronuz e Horus ERP.
- Validação e vínculo em tempo real de clientes (tabela CLIENTES) e fornecedores (tabela FORNECEDORES) do Horus.
- Armazenamento puro de códigos numéricos (sem prefixo 'H').
"""

import io
import csv
import os
import unicodedata
from pathlib import Path
from datetime import datetime, timezone
import logging
import re
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File, Form
from fastapi.responses import StreamingResponse
import openpyxl
from sqlalchemy.orm import Session
from sqlalchemy import or_, func, cast, String
from pydantic import BaseModel

from app.db.session import get_db
from app.core.dependencies import get_current_user
from app.core.upload_security import validate_file_size_and_extension, read_file_safely, sanitize_filename
from app.models.user import User, UserRole
from app.models.customer import Customer
from app.models.company_settings import CompanySettings
from app.integrators.horus_sql_client import HorusSQLClient, HorusSQLConfigError

router = APIRouter(prefix="/dbm", tags=["dbm"])
log = logging.getLogger(__name__)


def _resolve_seller_id(current_user: User, company_id: Optional[int] = None) -> int:
    is_master = False
    if hasattr(current_user, "type"):
        user_type = current_user.type
        type_str = user_type.value if hasattr(user_type, "value") else str(user_type or "")
        is_master = type_str.upper() == "MASTER"
    elif isinstance(current_user, dict):
        is_master = str(current_user.get("type", "")).upper() == "MASTER"

    if is_master and company_id and company_id > 0:
        return company_id

    seller_id = None
    if hasattr(current_user, "company_id"):
        seller_id = current_user.company_id
    elif isinstance(current_user, dict):
        seller_id = current_user.get("company_id")

    if not seller_id and company_id and company_id > 0:
        seller_id = company_id

    if not seller_id:
        raise HTTPException(
            status_code=400,
            detail="Seller não identificado. É necessário estar vinculado a uma empresa para gerenciar empresas no DBM."
        )
    return seller_id



def _sanitize_doc(doc: str) -> str:
    """Remove pontuação e espaços de CPF/CNPJ."""
    return re.sub(r"\D", "", doc or "")


def _normalize_header(h: str) -> str:
    if not h:
        return ""
    nfkd = unicodedata.normalize('NFKD', str(h).strip().lower())
    clean = re.sub(r'[^a-z0-9]', '', nfkd)
    return clean


def _map_columns(headers: List[str]) -> dict:
    col_map = {}
    for idx, raw_h in enumerate(headers):
        norm = _normalize_header(raw_h)
        if not norm:
            continue
        
        # codigo_horus
        if any(k in norm for k in ['codigohorus', 'codhorus', 'codcli', 'codigocliente', 'clientehorus', 'horus', 'codclient']) or norm in ['codigo', 'cod']:
            if 'codigo_horus' not in col_map:
                col_map['codigo_horus'] = idx
        # apelido / fantasia
        elif any(k in norm for k in ['apelido', 'fantasia', 'nomefantasia']) or norm in ['nome', 'empresa']:
            if 'apelido' not in col_map:
                col_map['apelido'] = idx
        # grupo
        elif any(k in norm for k in ['grupo', 'group']):
            if 'grupo' not in col_map:
                col_map['grupo'] = idx
        # segmento
        elif any(k in norm for k in ['segmento', 'segment', 'categoria']):
            if 'segmento' not in col_map:
                col_map['segmento'] = idx
        # obs da empresa
        elif any(k in norm for k in ['obs', 'observa', 'mensagem', 'msg', 'nota']):
            if 'obs_empresa' not in col_map:
                col_map['obs_empresa'] = idx
        # detalhes fechamento roy
        elif any(k in norm for k in ['royalt', 'fechamento']):
            if 'fechamento_roy' not in col_map:
                col_map['fechamento_roy'] = idx
        # customer account
        elif any(k in norm for k in ['account', 'conta', 'custummer']):
            if 'account' not in col_map:
                col_map['account'] = idx

    return col_map


# ── Schemas ───────────────────────────────────────────────────────────────────

class DbmImportItemPayload(BaseModel):
    row_index: int
    codigo_horus: int
    apelido: Optional[str] = None
    grupo: Optional[str] = None
    segmento: Optional[str] = None
    obs_empresa: Optional[str] = None
    fechamento_roy: Optional[str] = None
    account: Optional[str] = None


class DbmImportBatchPayload(BaseModel):
    company_id: Optional[int] = None
    seller_company_id: Optional[int] = None
    items: List[DbmImportItemPayload]

class DbmCompanySavePayload(BaseModel):
    id: Optional[int] = None
    seller_company_id: Optional[int] = None
    document: str
    name: str  # Apelido
    razao_social: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    group_name: Optional[str] = None
    segment: Optional[str] = None
    notes_message: Optional[str] = None
    customer_account: Optional[str] = None
    royalties_data: Optional[str] = None
    is_cliente: bool = True
    is_fornecedor: bool = False
    horus_cod_cli: Optional[int] = None
    horus_cod_fornecedor: Optional[int] = None


class DbmCompanyItem(BaseModel):
    id: int
    document: str
    name: str
    razao_social: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    group_name: Optional[str] = None
    segment: Optional[str] = None
    notes_message: Optional[str] = None
    customer_account: Optional[str] = None
    royalties_data: Optional[str] = None
    is_cliente: bool = True
    is_fornecedor: bool = False
    horus_cod_cli: Optional[int] = None
    horus_cod_fornecedor: Optional[int] = None
    created_at: Optional[str] = None


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("/companies/horus-lookup")
def lookup_horus_company(
    company_id: Optional[int] = Query(None, description="ID do seller para conexão Horus SQL"),
    document: str = Query(..., description="CNPJ ou CPF da empresa a pesquisar"),
    check_cliente: bool = Query(True, description="Consultar tabela CLIENTES do Horus"),
    check_fornecedor: bool = Query(True, description="Consultar tabela FORNECEDORES do Horus"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Pesquisa no banco SQL Server do Horus ERP por CPF/CNPJ:
    - Busca na tabela geral CLIENTES por COD_CLI, Razao Social, Apelido, Cidade e UF.
    - Busca na tabela FORNECEDORES amarrada ao cod_filial configurado nos dados da API (horus_branch).
    - Retorna dados completos para autopreenchimento imediato no formulário.
    """
    clean_doc = _sanitize_doc(document)
    if not clean_doc:
        raise HTTPException(status_code=400, detail="CNPJ/CPF inválido ou vazio para consulta no Horus.")

    # Identifica o seller / empresa ativo
    target_company_id = _resolve_seller_id(current_user, company_id)

    # Identificar CompanySettings correspondente
    settings = None
    if target_company_id:
        settings = db.query(CompanySettings).filter(CompanySettings.company_id == target_company_id).first()

    # Empresa para a conexão física Horus SQL:
    sql_company_id = target_company_id
    if not (settings and settings.horus_sql_enabled):
        active_sql = db.query(CompanySettings).filter(CompanySettings.horus_sql_enabled == True).first()
        if active_sql:
            sql_company_id = active_sql.company_id
            if not settings:
                settings = active_sql

    if not sql_company_id:
        raise HTTPException(status_code=400, detail="Nenhuma conexão Horus SQL configurada.")

    try:
        sql_client = HorusSQLClient(db, sql_company_id)
    except HorusSQLConfigError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        log.error(f"[DBM Horus Lookup] Falha ao inicializar client SQL: {e}")
        raise HTTPException(status_code=500, detail=f"Erro ao conectar ao banco Horus: {str(e)}")

    # ── Código de filial configurado nos dados da API (horus_branch) ──────────
    api_filial = None
    if settings:
        if settings.horus_branch and str(settings.horus_branch).strip():
            api_filial = str(settings.horus_branch).strip()
        elif settings.horus_sql_cod_filial and str(settings.horus_sql_cod_filial).strip():
            api_filial = str(settings.horus_sql_cod_filial).strip()

    cod_filial = int(api_filial) if (api_filial and api_filial.isdigit()) else (int(sql_client.cod_filial) if str(sql_client.cod_filial).isdigit() else 2)

    result = {
        "success": True,
        "clean_document": clean_doc,
        "filial_used": cod_filial,
        "cliente": {
            "found": False,
            "cod_cli": None,
            "nome_cli": None,
            "nome_reduzido": None,
            "cnpj_cpf": None,
            "cidade": None,
            "uf": None,
        },
        "fornecedor": {
            "found": False,
            "cod_fornecedor": None,
            "nom_fornecedor": None,
            "nom_fantasia": None,
            "cnpj_cpf": None,
            "cidade": None,
            "uf": None,
        },
    }

    # 1. Consulta em CLIENTES (Tabela geral) com Cidade e UF
    if check_cliente:
        try:
            sql_cli = """
                SELECT TOP 1 
                    c.COD_CLI, 
                    c.NOM_CLI, 
                    c.NOM_REDUZIDO, 
                    c.CNPJ, 
                    c.CPF,
                    u.SIGLA_UF,
                    cep.LOC_CIDADE
                FROM CLIENTES c WITH (NOLOCK)
                LEFT JOIN ENDERECOS_CLIENTE ec WITH (NOLOCK) 
                    ON ec.COD_CLI = c.COD_CLI
                LEFT JOIN TABELA_UF u WITH (NOLOCK) 
                    ON u.COD_UF = ec.COD_UF
                LEFT JOIN TABELA_CEP_IBGE cep WITH (NOLOCK) 
                    ON cep.COD_IBGE = ec.COD_IBGE
                WHERE REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(c.CNPJ, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                   OR REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(c.CPF, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                   OR c.CNPJ = %s
                   OR c.CPF = %s
                ORDER BY CASE WHEN ec.STA_DEFAULT = 'S' THEN 0 ELSE 1 END, ec.COD_END ASC
            """
            rows = sql_client.query(sql_cli, (clean_doc, clean_doc, clean_doc, clean_doc), max_rows=1)
            if rows and len(rows) > 0:
                row = rows[0]
                raw_cod = row.get("COD_CLI")
                cod_int = int(raw_cod) if str(raw_cod).isdigit() else None
                result["cliente"] = {
                    "found": True,
                    "cod_cli": cod_int,
                    "nome_cli": str(row.get("NOM_CLI") or "").strip(),
                    "nome_reduzido": str(row.get("NOM_REDUZIDO") or "").strip(),
                    "cnpj_cpf": str(row.get("CNPJ") or row.get("CPF") or "").strip(),
                    "cidade": str(row.get("LOC_CIDADE") or "").strip(),
                    "uf": str(row.get("SIGLA_UF") or "").strip().upper(),
                }
        except Exception as e:
            log.warning(f"[DBM Horus Lookup] Erro ao consultar CLIENTES: {e}")

    # 2. Consulta em FORNECEDORES (Tabela amarrada a COD_FILIAL da API)
    if check_fornecedor:
        try:
            sql_forn = f"""
                SELECT TOP 1 
                    f.COD_FORNECEDOR, 
                    f.NOM_FORNECEDOR, 
                    f.NOM_FANTASIA, 
                    f.CNPF_CNPJ, 
                    f.CPF,
                    u.SIGLA_UF,
                    cep.LOC_CIDADE
                FROM FORNECEDORES f WITH (NOLOCK)
                LEFT JOIN TABELA_UF u WITH (NOLOCK) 
                    ON u.COD_UF = f.UF
                LEFT JOIN TABELA_CEP_IBGE cep WITH (NOLOCK) 
                    ON cep.COD_IBGE = f.COD_IBGE
                WHERE f.COD_FILIAL = {cod_filial}
                  AND (
                    REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(f.CNPF_CNPJ, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                    OR REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(f.CPF, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                    OR f.CNPF_CNPJ = %s
                    OR f.CPF = %s
                  )
            """
            rows_forn = sql_client.query(sql_forn, (clean_doc, clean_doc, clean_doc, clean_doc), max_rows=1)
            if rows_forn and len(rows_forn) > 0:
                r_forn = rows_forn[0]
                raw_cod_forn = r_forn.get("COD_FORNECEDOR")
                cod_forn_int = int(raw_cod_forn) if str(raw_cod_forn).isdigit() else None
                result["fornecedor"] = {
                    "found": True,
                    "cod_fornecedor": cod_forn_int,
                    "nom_fornecedor": str(r_forn.get("NOM_FORNECEDOR") or "").strip(),
                    "nom_fantasia": str(r_forn.get("NOM_FANTASIA") or "").strip(),
                    "cnpj_cpf": str(r_forn.get("CNPF_CNPJ") or r_forn.get("CPF") or "").strip(),
                    "cidade": str(r_forn.get("LOC_CIDADE") or "").strip(),
                    "uf": str(r_forn.get("SIGLA_UF") or "").strip().upper(),
                }
            elif result["cliente"]["found"] and result["cliente"]["nome_cli"]:
                # Caso a linha do fornecedor esteja cadastrada na filial sem o CNPF_CNPJ preenchido,
                # localiza pelo nome do fornecedor/fantasia da filial
                cli_nome = result["cliente"]["nome_cli"]
                sql_forn_by_name = f"""
                    SELECT TOP 1 
                        f.COD_FORNECEDOR, 
                        f.NOM_FORNECEDOR, 
                        f.NOM_FANTASIA, 
                        f.CNPF_CNPJ, 
                        f.CPF,
                        u.SIGLA_UF,
                        cep.LOC_CIDADE
                    FROM FORNECEDORES f WITH (NOLOCK)
                    LEFT JOIN TABELA_UF u WITH (NOLOCK) 
                        ON u.COD_UF = f.UF
                    LEFT JOIN TABELA_CEP_IBGE cep WITH (NOLOCK) 
                        ON cep.COD_IBGE = f.COD_IBGE
                    WHERE f.COD_FILIAL = {cod_filial}
                      AND (
                        f.NOM_FORNECEDOR LIKE %s
                        OR %s LIKE f.NOM_FORNECEDOR + '%%'
                        OR f.NOM_FANTASIA LIKE %s
                      )
                """
                nome_param = f"%{cli_nome[:10]}%"
                rows_forn_name = sql_client.query(sql_forn_by_name, (nome_param, cli_nome, nome_param), max_rows=1)
                if rows_forn_name and len(rows_forn_name) > 0:
                    r_forn = rows_forn_name[0]
                    raw_cod_forn = r_forn.get("COD_FORNECEDOR")
                    cod_forn_int = int(raw_cod_forn) if str(raw_cod_forn).isdigit() else None
                    result["fornecedor"] = {
                        "found": True,
                        "cod_fornecedor": cod_forn_int,
                        "nom_fornecedor": str(r_forn.get("NOM_FORNECEDOR") or "").strip(),
                        "nom_fantasia": str(r_forn.get("NOM_FANTASIA") or "").strip(),
                        "cnpj_cpf": str(r_forn.get("CNPF_CNPJ") or r_forn.get("CPF") or "").strip(),
                        "cidade": str(r_forn.get("LOC_CIDADE") or "").strip(),
                        "uf": str(r_forn.get("SIGLA_UF") or "").strip().upper(),
                    }
        except Exception as e:
            log.warning(f"[DBM Horus Lookup] Erro ao consultar FORNECEDORES: {e}")

    return result


@router.get("/companies", response_model=dict)
def list_dbm_companies(
    company_id: Optional[int] = Query(None, description="Seller ativo para filtro contextual"),
    search: Optional[str] = Query(None, description="Busca por Fantasia (apelido), Razão Social, CNPJ ou Códigos Horus"),
    uf: Optional[str] = Query(None, description="Filtro por UF"),
    limit: int = Query(100, ge=1, le=500),
    skip: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Lista e busca empresas do Seller (tabela crm_customer) para a tela DBM.
    Garante isolamento absoluto por seller_company_id (sem expor outras empresas/tenants do sistema).
    """
    seller_id = _resolve_seller_id(current_user, company_id)

    query = db.query(Customer).filter(Customer.company_id == seller_id)

    if uf and uf.strip():
        query = query.filter(func.upper(Customer.state) == uf.strip().upper())

    if search and search.strip():
        term = f"%{search.strip()}%"
        clean_term = _sanitize_doc(search.strip())

        filters = [
            Customer.name.ilike(term),
            Customer.corporate_name.ilike(term),
            Customer.document.ilike(term),
        ]
        if clean_term:
            filters.append(func.regexp_replace(Customer.document, r"\D", "", "g").ilike(f"%{clean_term}%"))
            if clean_term.isdigit():
                clean_num = int(clean_term)
                filters.append(Customer.id == clean_num)
                filters.append(Customer.horus_cod_cli == clean_num)
                filters.append(Customer.horus_cod_fornecedor == clean_num)

        query = query.filter(or_(*filters))

    total = query.count()
    customers = query.order_by(Customer.id.desc()).offset(skip).limit(limit).all()

    items = []
    for c in customers:
        items.append({
            "id": c.id,
            "document": c.document or "",
            "name": c.name or "",
            "razao_social": c.corporate_name or c.name or "",
            "city": c.city or "",
            "state": c.state or "",
            "group_name": c.group_name or "",
            "segment": c.segment or "",
            "notes_message": c.notes_message or "",
            "customer_account": c.customer_account or "",
            "royalties_data": c.royalties_data or "",
            "is_cliente": c.is_cliente if c.is_cliente is not None else True,
            "is_fornecedor": c.is_fornecedor if c.is_fornecedor is not None else False,
            "horus_cod_cli": c.horus_cod_cli,
            "horus_cod_fornecedor": c.horus_cod_fornecedor,
            "created_at": c.created_at.isoformat() if c.created_at else None,
        })

    return {
        "items": items,
        "total": total,
        "skip": skip,
        "limit": limit,
    }


@router.get("/companies/{id}")
def get_dbm_company(
    id: int,
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Retorna os dados completos da empresa do seller pelo Código (Customer ID).
    """
    seller_id = _resolve_seller_id(current_user, company_id)
    cust = db.query(Customer).filter(
        Customer.id == id,
        Customer.company_id == seller_id
    ).first()
    if not cust:
        raise HTTPException(status_code=404, detail="Empresa não encontrada para este seller.")

    return {
        "id": cust.id,
        "document": cust.document or "",
        "name": cust.name or "",
        "razao_social": cust.corporate_name or cust.name or "",
        "city": cust.city or "",
        "state": cust.state or "",
        "group_name": cust.group_name or "",
        "segment": cust.segment or "",
        "notes_message": cust.notes_message or "",
        "customer_account": cust.customer_account or "",
        "royalties_data": cust.royalties_data or "",
        "is_cliente": cust.is_cliente if cust.is_cliente is not None else True,
        "is_fornecedor": cust.is_fornecedor if cust.is_fornecedor is not None else False,
        "horus_cod_cli": cust.horus_cod_cli,
        "horus_cod_fornecedor": cust.horus_cod_fornecedor,
        "created_at": cust.created_at.isoformat() if cust.created_at else None,
    }


@router.post("/companies")
def save_dbm_company(
    payload: DbmCompanySavePayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Cadastra ou atualiza uma empresa do seller (crm_customer):
    - Valida se já existe na tabela crm_customer do seller (por ID ou por CNPJ/CPF).
    - Se já existir, atualiza os dados e vínculos com Horus.
    - Se não existir, cadastra isolado na conta deste seller.
    - NUNCA altera ou polui a tabela cmp_company (SaaS tenants).
    """
    seller_id = _resolve_seller_id(current_user, payload.seller_company_id)

    clean_doc = _sanitize_doc(payload.document)
    if not clean_doc:
        raise HTTPException(status_code=400, detail="CNPJ / CPF é obrigatório para cadastrar a empresa.")

    apelido = (payload.name or "").strip()
    if not apelido:
        raise HTTPException(status_code=400, detail="Fantasia (apelido) é obrigatório.")

    razao_social = (payload.razao_social or apelido).strip()

    # Validação de existência no escopo do seller
    existing_cust = None
    if payload.id and payload.id > 0:
        existing_cust = db.query(Customer).filter(
            Customer.id == payload.id,
            Customer.company_id == seller_id
        ).first()

    if not existing_cust:
        existing_cust = db.query(Customer).filter(
            Customer.company_id == seller_id,
            or_(
                func.regexp_replace(Customer.document, r"\D", "", "g") == clean_doc,
                Customer.document == payload.document.strip()
            )
        ).first()

    status_action = "updated" if existing_cust else "created"

    if existing_cust:
        existing_cust.name = apelido
        existing_cust.corporate_name = razao_social
        existing_cust.document = clean_doc
        existing_cust.city = (payload.city or "").strip() or None
        existing_cust.state = (payload.state or "").strip().upper() or None
        existing_cust.group_name = (payload.group_name or "").strip() or None
        existing_cust.segment = (payload.segment or "").strip() or None
        existing_cust.notes_message = (payload.notes_message or "").strip() or None
        existing_cust.customer_account = (payload.customer_account or "").strip() or None
        existing_cust.royalties_data = (payload.royalties_data or "").strip() or None
        existing_cust.is_cliente = payload.is_cliente
        existing_cust.is_fornecedor = payload.is_fornecedor
        existing_cust.horus_cod_cli = payload.horus_cod_cli
        existing_cust.horus_cod_fornecedor = payload.horus_cod_fornecedor
        target_cust = existing_cust
    else:
        new_cust = Customer(
            company_id=seller_id,
            name=apelido,
            corporate_name=razao_social,
            document=clean_doc,
            city=(payload.city or "").strip() or None,
            state=(payload.state or "").strip().upper() or None,
            group_name=(payload.group_name or "").strip() or None,
            segment=(payload.segment or "").strip() or None,
            notes_message=(payload.notes_message or "").strip() or None,
            customer_account=(payload.customer_account or "").strip() or None,
            royalties_data=(payload.royalties_data or "").strip() or None,
            is_cliente=payload.is_cliente,
            is_fornecedor=payload.is_fornecedor,
            horus_cod_cli=payload.horus_cod_cli,
            horus_cod_fornecedor=payload.horus_cod_fornecedor,
            customer_type="PJ" if len(clean_doc) > 11 else "PF",
        )
        db.add(new_cust)
        db.flush()
        target_cust = new_cust

    db.commit()
    db.refresh(target_cust)

    return {
        "success": True,
        "action": status_action,
        "message": f"Empresa {status_action} com sucesso no Cronuz!",
        "company": {
            "id": target_cust.id,
            "document": target_cust.document,
            "name": target_cust.name,
            "razao_social": target_cust.corporate_name,
            "city": target_cust.city,
            "state": target_cust.state,
            "group_name": target_cust.group_name,
            "segment": target_cust.segment,
            "notes_message": target_cust.notes_message,
            "customer_account": target_cust.customer_account,
            "royalties_data": target_cust.royalties_data,
            "is_cliente": target_cust.is_cliente,
            "is_fornecedor": target_cust.is_fornecedor,
            "horus_cod_cli": target_cust.horus_cod_cli,
            "horus_cod_fornecedor": target_cust.horus_cod_fornecedor,
            "created_at": target_cust.created_at.isoformat() if target_cust.created_at else None,
        }
    }


# ── Importação Inteligente de Planilha Excel / CSV ────────────────────────────

@router.get("/companies/import/template")
def download_dbm_import_template():
    """
    Gera e entrega modelo Excel (.xlsx) para importação de empresas no DBM.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Importacao DBM"

    headers = [
        "codigo_horus",
        "apelido",
        "grupo",
        "segmento",
        "obs da empresa",
        "detalhes fechamento roy",
        "account"
    ]
    ws.append(headers)

    sample_rows = [
        [11740, "EO EDITORA", "Grupo Editorial", "Editoras", "Instruções de faturamento e notas", "Banco do Brasil Ag 1234 CC 5678", "ACC-1001"],
        [2, "Livraria Cultura", "Varejo", "Livrarias", "Entregar somente em horário comercial", "", "ACC-2002"]
    ]
    for r in sample_rows:
        ws.append(r)

    # Ajusta largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "modelo_importacao_empresas_dbm.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/companies/import/parse")
async def parse_dbm_companies_sheet(
    file: UploadFile = File(...),
    company_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Recebe planilha (.xlsx ou .csv), valida integridade e tamanho (máx 10 MB),
    faz parsing flexível das colunas e retorna as linhas com código Horus.
    """
    clean_filename = validate_file_size_and_extension(file, category="sheet")
    file_bytes = await read_file_safely(file, max_size_bytes=10 * 1024 * 1024)

    # Identificar seller / isolamento de pastas
    target_seller_id = _resolve_seller_id(current_user, company_id)

    # Salva na pasta isolada de uploads do seller (uploads/<company_id>/sheets/)
    upload_dir = Path("uploads") / str(target_seller_id) / "sheets"
    upload_dir.mkdir(parents=True, exist_ok=True)
    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    dest_path = upload_dir / f"dbm_import_{timestamp_str}_{clean_filename}"
    try:
        with open(dest_path, "wb") as f:
            f.write(file_bytes)
    except Exception as e:
        log.warning(f"[DBM Import] Não foi possível persistir cópia da planilha em disco: {e}")

    # Leitura das linhas
    ext = Path(clean_filename).suffix.lower()
    raw_rows = []

    if ext in [".xlsx", ".ods"]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(values_only=True):
                cleaned_row = ["" if v is None else str(v).strip() for v in row]
                if any(cleaned_row):
                    raw_rows.append(cleaned_row)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Erro ao ler arquivo Excel (.xlsx): {str(e)}")
    elif ext == ".csv":
        try:
            text_content = None
            for encoding in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
                try:
                    text_content = file_bytes.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if not text_content:
                raise HTTPException(status_code=400, detail="Não foi possível decodificar o arquivo CSV.")
            
            sample = text_content[:2048]
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
            reader = csv.reader(io.StringIO(text_content), delimiter=delimiter)
            for r in reader:
                cleaned_row = [str(c).strip() for c in r]
                if any(cleaned_row):
                    raw_rows.append(cleaned_row)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Erro ao ler arquivo CSV: {str(e)}")

    if not raw_rows or len(raw_rows) < 2:
        raise HTTPException(status_code=400, detail="A planilha está vazia ou não contém cabeçalhos e registros válidos.")

    headers = raw_rows[0]
    col_map = _map_columns(headers)

    if "codigo_horus" not in col_map:
        raise HTTPException(
            status_code=400, 
            detail="Não foi possível identificar a coluna com o Código do Horus. Verifique se o cabeçalho contém 'codigo_horus', 'cod_horus' ou 'cod_cli'."
        )

    parsed_items = []
    idx_cod = col_map["codigo_horus"]
    idx_apelido = col_map.get("apelido")
    idx_grupo = col_map.get("grupo")
    idx_seg = col_map.get("segmento")
    idx_obs = col_map.get("obs_empresa")
    idx_roy = col_map.get("fechamento_roy")
    idx_acc = col_map.get("account")

    for row_idx, r in enumerate(raw_rows[1:], start=2):
        if idx_cod >= len(r):
            continue
        raw_val = r[idx_cod].strip()
        raw_val_clean = re.sub(r"[^\d]", "", raw_val)
        if not raw_val_clean.isdigit():
            continue
        
        cod_horus_int = int(raw_val_clean)
        if cod_horus_int <= 0:
            continue

        apelido = r[idx_apelido].strip() if idx_apelido is not None and idx_apelido < len(r) else ""
        grupo = r[idx_grupo].strip() if idx_grupo is not None and idx_grupo < len(r) else ""
        segmento = r[idx_seg].strip() if idx_seg is not None and idx_seg < len(r) else ""
        obs_empresa = r[idx_obs].strip() if idx_obs is not None and idx_obs < len(r) else ""
        fechamento_roy = r[idx_roy].strip() if idx_roy is not None and idx_roy < len(r) else ""
        account = r[idx_acc].strip() if idx_acc is not None and idx_acc < len(r) else ""

        parsed_items.append({
            "row_index": row_idx,
            "codigo_horus": cod_horus_int,
            "apelido": apelido or None,
            "grupo": grupo or None,
            "segmento": segmento or "Editoras",
            "obs_empresa": obs_empresa or None,
            "fechamento_roy": fechamento_roy or None,
            "account": account or None,
        })

    if not parsed_items:
        raise HTTPException(
            status_code=400,
            detail="Nenhuma linha com código de cliente Horus válido foi encontrada na planilha."
        )

    return {
        "success": True,
        "filename": clean_filename,
        "total_rows": len(parsed_items),
        "columns_detected": list(col_map.keys()),
        "preview": parsed_items[:5],
        "rows": parsed_items
    }


@router.post("/companies/import/process-batch")
def process_dbm_companies_import_batch(
    payload: DbmImportBatchPayload,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Processa um lote de empresas parseadas da planilha para o Seller (crm_customer):
    1. Para cada empresa, busca dados no Horus pelo codigo_horus (tabela CLIENTES + FORNECEDORES).
    2. Preserva os dados informados na planilha (apelido, grupo, segmento, obs, royalties, account).
    3. Cadastra ou atualiza na tabela crm_customer do Seller.
    4. NUNCA toca na tabela cmp_company (SaaS tenants).
    """
    seller_id = _resolve_seller_id(current_user, payload.seller_company_id or payload.company_id)

    # Identificar CompanySettings correspondente
    settings = db.query(CompanySettings).filter(CompanySettings.company_id == seller_id).first()

    # Empresa de conexão Horus SQL
    sql_company_id = seller_id
    if not (settings and settings.horus_sql_enabled):
        active_sql = db.query(CompanySettings).filter(CompanySettings.horus_sql_enabled == True).first()
        if active_sql:
            sql_company_id = active_sql.company_id
            if not settings:
                settings = active_sql

    if not sql_company_id:
        raise HTTPException(status_code=400, detail="Nenhuma conexão Horus SQL configurada no sistema.")

    try:
        sql_client = HorusSQLClient(db, sql_company_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao conectar ao banco Horus: {str(e)}")

    # Filial configurada nos dados da API (horus_branch)
    api_filial = None
    if settings:
        if settings.horus_branch and str(settings.horus_branch).strip():
            api_filial = str(settings.horus_branch).strip()
        elif settings.horus_sql_cod_filial and str(settings.horus_sql_cod_filial).strip():
            api_filial = str(settings.horus_sql_cod_filial).strip()

    cod_filial = int(api_filial) if (api_filial and api_filial.isdigit()) else (int(sql_client.cod_filial) if str(sql_client.cod_filial).isdigit() else 2)

    results = []
    
    for item in payload.items:
        cod_cli = item.codigo_horus
        row_idx = item.row_index

        try:
            # 1. Consulta no Horus CLIENTES
            sql_cli = """
                SELECT TOP 1 
                    c.COD_CLI, c.NOM_CLI, c.NOM_REDUZIDO, c.CNPJ, c.CPF,
                    u.SIGLA_UF, cep.LOC_CIDADE
                FROM CLIENTES c WITH (NOLOCK)
                LEFT JOIN ENDERECOS_CLIENTE ec WITH (NOLOCK) ON ec.COD_CLI = c.COD_CLI
                LEFT JOIN TABELA_UF u WITH (NOLOCK) ON u.COD_UF = ec.COD_UF
                LEFT JOIN TABELA_CEP_IBGE cep WITH (NOLOCK) ON cep.COD_IBGE = ec.COD_IBGE
                WHERE c.COD_CLI = %s
                ORDER BY CASE WHEN ec.STA_DEFAULT = 'S' THEN 0 ELSE 1 END, ec.COD_END ASC
            """
            rows = sql_client.query(sql_cli, (cod_cli,), max_rows=1)
            
            if not rows or len(rows) == 0:
                results.append({
                    "row_index": row_idx,
                    "codigo_horus": cod_cli,
                    "apelido": item.apelido or f"Horus #{cod_cli}",
                    "status": "WARNING",
                    "action": "SKIPPED",
                    "message": f"Cliente #{cod_cli} não localizado na tabela CLIENTES do Horus ERP."
                })
                continue

            r_cli = rows[0]
            nom_cli = str(r_cli.get("NOM_CLI") or "").strip()
            nom_reduzido = str(r_cli.get("NOM_REDUZIDO") or "").strip()
            raw_doc = str(r_cli.get("CNPJ") or r_cli.get("CPF") or "").strip()
            clean_doc = _sanitize_doc(raw_doc)
            city_horus = str(r_cli.get("LOC_CIDADE") or "").strip() or None
            uf_horus = str(r_cli.get("SIGLA_UF") or "").strip().upper() or None

            # Fallback se não tiver CNPJ nem CPF no Horus
            if not clean_doc:
                clean_doc = f"HORUS{cod_cli:08d}"

            # 2. Consulta no Horus FORNECEDORES na filial da API
            horus_cod_fornecedor = None
            is_fornecedor = False
            try:
                sql_forn = f"""
                    SELECT TOP 1 f.COD_FORNECEDOR
                    FROM FORNECEDORES f WITH (NOLOCK)
                    WHERE f.COD_FILIAL = {cod_filial}
                      AND (
                        REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(f.CNPF_CNPJ, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                        OR REPLACE(REPLACE(REPLACE(REPLACE(ISNULL(f.CPF, ''), '.', ''), '/', ''), '-', ''), ' ', '') = %s
                        OR f.NOM_FORNECEDOR = %s
                        OR f.NOM_FANTASIA = %s
                      )
                """
                rows_forn = sql_client.query(sql_forn, (clean_doc, clean_doc, nom_cli, nom_reduzido or nom_cli), max_rows=1)
                if rows_forn and len(rows_forn) > 0:
                    raw_forn = rows_forn[0].get("COD_FORNECEDOR")
                    if str(raw_forn).isdigit():
                        horus_cod_fornecedor = int(raw_forn)
                        is_fornecedor = True
            except Exception as ef:
                log.warning(f"[DBM Batch Import] Erro ao buscar fornecedor #{cod_cli}: {ef}")

            # 3. Mesclagem inteligente de dados
            # A planilha tem PRIORIDADE para: apelido, grupo, segmento, obs, royalties, account
            apelido = (item.apelido or "").strip()
            if not apelido:
                apelido = nom_reduzido or nom_cli or f"Cliente {cod_cli}"

            razao_social = nom_cli or apelido
            grupo = (item.grupo or "").strip() or None
            segmento = (item.segmento or "").strip() or "Editoras"
            obs = (item.obs_empresa or "").strip() or None
            fechamento_roy = (item.fechamento_roy or "").strip() or None
            account = (item.account or "").strip() or None

            # 4. Localização / Atualização / Criação no crm_customer DO SELLER
            existing_cust = db.query(Customer).filter(
                Customer.company_id == seller_id,
                or_(
                    Customer.horus_cod_cli == cod_cli,
                    func.regexp_replace(Customer.document, r"\D", "", "g") == clean_doc,
                    Customer.document == clean_doc
                )
            ).first()

            action = "UPDATED" if existing_cust else "CREATED"

            if existing_cust:
                existing_cust.name = apelido
                if nom_cli:
                    existing_cust.corporate_name = nom_cli
                if clean_doc:
                    existing_cust.document = clean_doc
                if city_horus:
                    existing_cust.city = city_horus
                if uf_horus:
                    existing_cust.state = uf_horus
                if grupo:
                    existing_cust.group_name = grupo
                if segmento:
                    existing_cust.segment = segmento
                if obs:
                    existing_cust.notes_message = obs
                if account:
                    existing_cust.customer_account = account
                if fechamento_roy:
                    existing_cust.royalties_data = fechamento_roy

                existing_cust.is_cliente = True
                existing_cust.horus_cod_cli = cod_cli
                if is_fornecedor and horus_cod_fornecedor:
                    existing_cust.is_fornecedor = True
                    existing_cust.horus_cod_fornecedor = horus_cod_fornecedor

                target_cust = existing_cust
            else:
                new_cust = Customer(
                    company_id=seller_id,
                    document=clean_doc,
                    name=apelido,
                    corporate_name=razao_social,
                    city=city_horus,
                    state=uf_horus,
                    group_name=grupo,
                    segment=segmento,
                    notes_message=obs,
                    customer_account=account,
                    royalties_data=fechamento_roy,
                    is_cliente=True,
                    is_fornecedor=is_fornecedor,
                    horus_cod_cli=cod_cli,
                    horus_cod_fornecedor=horus_cod_fornecedor,
                    customer_type="PJ" if len(clean_doc) > 11 else "PF",
                )
                db.add(new_cust)
                db.flush()
                target_cust = new_cust

            db.commit()

            results.append({
                "row_index": row_idx,
                "codigo_horus": cod_cli,
                "apelido": apelido,
                "razao_social": razao_social,
                "document": clean_doc,
                "company_id": target_cust.id,
                "action": action,
                "status": "SUCCESS",
                "message": f"Empresa '{apelido}' ({action.lower()}) vinculada ao Horus #{cod_cli}"
            })

        except Exception as e:
            db.rollback()
            log.error(f"[DBM Batch Import] Erro no item linha {row_idx} (Horus #{cod_cli}): {e}")
            results.append({
                "row_index": row_idx,
                "codigo_horus": cod_cli,
                "apelido": item.apelido or f"Horus #{cod_cli}",
                "status": "ERROR",
                "action": "FAILED",
                "message": f"Erro interno ao processar linha {row_idx}: {str(e)}"
            })

    return {
        "success": True,
        "seller_company_id": seller_id,
        "total_processed": len(results),
        "results": results
    }
