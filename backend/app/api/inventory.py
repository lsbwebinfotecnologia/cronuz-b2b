import io
import csv
import secrets
import logging
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, text, desc, case
import openpyxl

from app.db.session import get_db
from app.core import dependencies, security
from app.models import user as user_models
from app.models.company import Company
from app.models.inventory import (
    Inventory,
    InventoryItem,
    InventorySession,
    InventoryScan,
    InventoryStatus,
    SessionStatus,
    SessionType,
)
from app.schemas import inventory as inv_schemas

router = APIRouter(prefix="/companies/{company_id}/inventory", tags=["inventory"])
public_router = APIRouter(prefix="/inventory/public", tags=["inventory-public"])
logger = logging.getLogger("cronuz.inventory")


def _assert_inventory_access(current_user: user_models.User, company_id: int, db: Session) -> Company:
    """Valida se o usuário tem permissão para acessar o inventário da empresa."""
    user_type = getattr(current_user.type, "value", str(current_user.type))
    if user_type != "MASTER" and current_user.company_id != company_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Acesso restrito a esta empresa.")
    
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Empresa não encontrada.")
    
    if user_type != "MASTER" and not getattr(company, "has_inventory_module", False):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN, 
            detail="Módulo de Inventário não está ativo para esta empresa. Contate o administrador."
        )
    
    return company


def _get_inventory_by_token(access_token: str, db: Session) -> Inventory:
    """Busca o inventário pelo token público garantindo que o acesso esteja ativo."""
    inv = db.query(Inventory).filter(Inventory.access_token == access_token).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Link de inventário inválido ou inexistente.")
    if not getattr(inv, "is_public_access_enabled", True):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="O acesso rápido por link a este inventário foi desativado pelo gestor.")
    return inv


# ======================================================================
# ROTAS AUTENTICADAS DO GESTOR (ROUTER)
# ======================================================================

# ----------------------------------------------------------------------
# 1. Listagem e Criação de Inventários
# ----------------------------------------------------------------------
@router.get("", response_model=List[inv_schemas.InventoryResponse])
def list_inventories(
    company_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    inventories = (
        db.query(Inventory)
        .filter(Inventory.company_id == company_id)
        .order_by(desc(Inventory.created_at))
        .all()
    )
    
    result = []
    for inv in inventories:
        total_scanned = (
            db.query(func.coalesce(func.sum(InventoryScan.quantity), 0))
            .filter(InventoryScan.inventory_id == inv.id)
            .scalar()
        )
        total_sessions = db.query(func.count(InventorySession.id)).filter(InventorySession.inventory_id == inv.id).scalar()
        open_sessions = (
            db.query(func.count(InventorySession.id))
            .filter(InventorySession.inventory_id == inv.id, InventorySession.status == SessionStatus.ABERTA.value)
            .scalar()
        )
        
        # Garante token se inventário for antigo
        if not inv.access_token:
            inv.access_token = secrets.token_urlsafe(32)
            db.commit()

        inv_dict = {
            "id": inv.id,
            "company_id": inv.company_id,
            "code": inv.code,
            "name": inv.name,
            "status": inv.status,
            "description": inv.description,
            "total_expected_skus": inv.total_expected_skus,
            "total_scanned_items": int(total_scanned or 0),
            "total_sessions": int(total_sessions or 0),
            "open_sessions": int(open_sessions or 0),
            "access_token": inv.access_token,
            "is_public_access_enabled": inv.is_public_access_enabled,
            "created_by_user_id": inv.created_by_user_id,
            "finalized_by_user_id": inv.finalized_by_user_id,
            "finalized_at": inv.finalized_at,
            "created_at": inv.created_at,
            "updated_at": inv.updated_at,
        }
        result.append(inv_schemas.InventoryResponse(**inv_dict))
        
    return result


@router.post("", response_model=inv_schemas.InventoryResponse)
def create_inventory(
    company_id: int,
    payload: inv_schemas.InventoryCreate,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    code = payload.code
    if not code or not code.strip():
        now = datetime.now()
        count_year = (
            db.query(func.count(Inventory.id))
            .filter(Inventory.company_id == company_id)
            .scalar()
            or 0
        )
        code = f"INV-{now.year}-{str(count_year + 1).zfill(3)}"
    
    access_token = secrets.token_urlsafe(32)

    new_inv = Inventory(
        company_id=company_id,
        code=code.strip(),
        name=payload.name.strip(),
        description=payload.description.strip() if payload.description else None,
        status=InventoryStatus.EM_ANDAMENTO.value,
        total_expected_skus=0,
        access_token=access_token,
        is_public_access_enabled=True,
        supervisor_pin=payload.supervisor_pin.strip() if payload.supervisor_pin else "1234",
        created_by_user_id=current_user.id,
    )
    db.add(new_inv)
    db.commit()
    db.refresh(new_inv)
    
    return inv_schemas.InventoryResponse(
        id=new_inv.id,
        company_id=new_inv.company_id,
        code=new_inv.code,
        name=new_inv.name,
        status=new_inv.status,
        description=new_inv.description,
        total_expected_skus=new_inv.total_expected_skus,
        total_scanned_items=0,
        total_sessions=0,
        open_sessions=0,
        access_token=new_inv.access_token,
        is_public_access_enabled=new_inv.is_public_access_enabled,
        supervisor_pin=new_inv.supervisor_pin or "1234",
        created_by_user_id=new_inv.created_by_user_id,
        finalized_by_user_id=None,
        finalized_at=None,
        created_at=new_inv.created_at,
        updated_at=new_inv.updated_at,
    )


# ----------------------------------------------------------------------
# 2. Detalhes de um Inventário Específico
# ----------------------------------------------------------------------
@router.get("/{inventory_id}", response_model=inv_schemas.InventoryResponse)
def get_inventory(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    # Garante token se estiver vazio
    if not inv.access_token:
        inv.access_token = secrets.token_urlsafe(32)
        db.commit()

    total_scanned = (
        db.query(func.coalesce(func.sum(InventoryScan.quantity), 0))
        .filter(InventoryScan.inventory_id == inv.id)
        .scalar()
    )
    total_sessions = db.query(func.count(InventorySession.id)).filter(InventorySession.inventory_id == inv.id).scalar()
    open_sessions = (
        db.query(func.count(InventorySession.id))
        .filter(InventorySession.inventory_id == inv.id, InventorySession.status == SessionStatus.ABERTA.value)
        .scalar()
    )
    
    return inv_schemas.InventoryResponse(
        id=inv.id,
        company_id=inv.company_id,
        code=inv.code,
        name=inv.name,
        status=inv.status,
        description=inv.description,
        total_expected_skus=inv.total_expected_skus,
        total_scanned_items=int(total_scanned or 0),
        total_sessions=int(total_sessions or 0),
        open_sessions=int(open_sessions or 0),
        access_token=inv.access_token,
        is_public_access_enabled=inv.is_public_access_enabled,
        supervisor_pin=inv.supervisor_pin or "1234",
        created_by_user_id=inv.created_by_user_id,
        finalized_by_user_id=inv.finalized_by_user_id,
        finalized_at=inv.finalized_at,
        created_at=inv.created_at,
        updated_at=inv.updated_at,
    )


@router.post("/{inventory_id}/regenerate-token")
def regenerate_inventory_token(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    inv.access_token = secrets.token_urlsafe(32)
    db.commit()
    db.refresh(inv)
    return {"access_token": inv.access_token}


# ----------------------------------------------------------------------
# 3. Carga da Base com Validação Estrita de Duplicidade
# ----------------------------------------------------------------------
@router.post("/{inventory_id}/upload-sheet")
async def upload_inventory_sheet(
    company_id: int,
    inventory_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    if inv.status != InventoryStatus.EM_ANDAMENTO.value:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Operação bloqueada. O inventário não está em andamento (Status atual: '{inv.status}')."
        )
    
    filename = file.filename.lower()
    content = await file.read()
    
    raw_rows = []
    
    if filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        headers = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in sheet[1]]
        
        isbn_idx = next((i for i, h in enumerate(headers) if h in ["isbn", "ean", "codigo", "codigo de barras"]), None)
        title_idx = next((i for i, h in enumerate(headers) if h in ["titulo", "título", "nome", "descricao", "descrição"]), None)
        publisher_idx = next((i for i, h in enumerate(headers) if h in ["editora", "marca", "fabricante"]), None)
        category_idx = next((i for i, h in enumerate(headers) if h in ["categoria", "secao", "seção"]), None)
        location_idx = next((i for i, h in enumerate(headers) if h in ["endereco", "endereço", "local", "localizacao", "localização", "prateleira"]), None)
        
        if isbn_idx is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Coluna 'ISBN' ou 'Código de Barras' não encontrada na planilha.")
        
        for line_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue
            isbn_val = str(row[isbn_idx]).strip() if row[isbn_idx] is not None else ""
            if not isbn_val or isbn_val == "None":
                continue
            if isbn_val.endswith(".0"):
                isbn_val = isbn_val[:-2]
            
            title_val = str(row[title_idx]).strip() if title_idx is not None and row[title_idx] is not None else "Sem Título"
            pub_val = str(row[publisher_idx]).strip() if publisher_idx is not None and row[publisher_idx] is not None else None
            cat_val = str(row[category_idx]).strip() if category_idx is not None and row[category_idx] is not None else None
            loc_val = str(row[location_idx]).strip() if location_idx is not None and row[location_idx] is not None else None
            
            raw_rows.append({
                "line": line_num,
                "isbn": isbn_val,
                "title": title_val[:255],
                "publisher": pub_val[:255] if pub_val else None,
                "category": cat_val[:100] if cat_val else None,
                "default_location": loc_val[:100] if loc_val else None,
            })
            
    elif filename.endswith(".csv"):
        text_content = ""
        for encoding in ["utf-8-sig", "utf-8", "latin-1"]:
            try:
                text_content = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        
        if not text_content:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Codificação do arquivo CSV não suportada.")
            
        csv_reader = csv.reader(io.StringIO(text_content), delimiter=";" if ";" in text_content[:200] else ",")
        rows_list = list(csv_reader)
        if not rows_list:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo CSV vazio.")
            
        headers = [h.strip().lower() for h in rows_list[0]]
        isbn_idx = next((i for i, h in enumerate(headers) if h in ["isbn", "ean", "codigo", "codigo de barras"]), None)
        title_idx = next((i for i, h in enumerate(headers) if h in ["titulo", "título", "nome", "descricao", "descrição"]), None)
        publisher_idx = next((i for i, h in enumerate(headers) if h in ["editora", "marca", "fabricante"]), None)
        category_idx = next((i for i, h in enumerate(headers) if h in ["categoria", "secao", "seção"]), None)
        location_idx = next((i for i, h in enumerate(headers) if h in ["endereco", "endereço", "local", "localizacao", "localização", "prateleira"]), None)
        
        if isbn_idx is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Coluna 'ISBN' não encontrada no arquivo CSV.")
            
        for line_num, row in enumerate(rows_list[1:], start=2):
            if not row or len(row) <= isbn_idx:
                continue
            isbn_val = str(row[isbn_idx]).strip()
            if not isbn_val:
                continue
            title_val = str(row[title_idx]).strip() if title_idx is not None and len(row) > title_idx else "Sem Título"
            pub_val = str(row[publisher_idx]).strip() if publisher_idx is not None and len(row) > publisher_idx else None
            cat_val = str(row[category_idx]).strip() if category_idx is not None and len(row) > category_idx else None
            loc_val = str(row[location_idx]).strip() if location_idx is not None and len(row) > location_idx else None
            
            raw_rows.append({
                "line": line_num,
                "isbn": isbn_val,
                "title": title_val[:255],
                "publisher": pub_val[:255] if pub_val else None,
                "category": cat_val[:100] if cat_val else None,
                "default_location": loc_val[:100] if loc_val else None,
            })
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Formato de arquivo inválido. Envie um arquivo .xlsx ou .csv.")

    if not raw_rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nenhum item válido encontrado no arquivo.")

    # 1. Deduplica linhas da própria planilha mantendo a 1ª ocorrência de cada ISBN
    unique_raw_rows = []
    seen_isbns = set()
    for item in raw_rows:
        isbn = item["isbn"]
        if isbn not in seen_isbns:
            seen_isbns.add(isbn)
            unique_raw_rows.append(item)

    # 2. Busca itens que já existem no banco de dados para este inventário
    existing_items = (
        db.query(InventoryItem)
        .filter(InventoryItem.inventory_id == inventory_id, InventoryItem.isbn.in_(seen_isbns))
        .all()
    )
    existing_map = {item.isbn: item for item in existing_items}

    inserted_count = 0
    updated_unregistered_count = 0
    ignored_count = 0

    for it in unique_raw_rows:
        isbn = it["isbn"]
        if isbn in existing_map:
            db_item = existing_map[isbn]
            if db_item.is_unregistered:
                # Atualiza metadados do item que havia sido bipado como avulso previamente
                db_item.title = it["title"]
                db_item.publisher = it["publisher"]
                db_item.category = it["category"]
                db_item.default_location = it["default_location"]
                db_item.is_unregistered = False
                updated_unregistered_count += 1
            else:
                # Já cadastrado oficialmente: ignora conforme solicitado pelo usuário
                ignored_count += 1
        else:
            # Novo cadastro
            db.add(InventoryItem(
                inventory_id=inventory_id,
                company_id=company_id,
                isbn=isbn,
                title=it["title"],
                publisher=it["publisher"],
                category=it["category"],
                default_location=it["default_location"],
                is_unregistered=False
            ))
            inserted_count += 1

    db.commit()

    total_skus = db.query(func.count(InventoryItem.id)).filter(InventoryItem.inventory_id == inventory_id).scalar()
    inv.total_expected_skus = total_skus or 0
    db.commit()

    msg_parts = []
    if inserted_count > 0:
        msg_parts.append(f"{inserted_count} novo(s) produto(s) cadastrado(s)")
    if updated_unregistered_count > 0:
        msg_parts.append(f"{updated_unregistered_count} item(ns) avulso(s) atualizado(s)")
    if ignored_count > 0:
        msg_parts.append(f"{ignored_count} produto(s) já cadastrado(s) ignorado(s)")
    
    final_message = "; ".join(msg_parts) + "." if msg_parts else "Nenhum novo produto para cadastrar."

    return {
        "message": final_message,
        "inserted_count": inserted_count,
        "ignored_count": ignored_count,
        "updated_unregistered_count": updated_unregistered_count,
        "total_expected_skus": inv.total_expected_skus
    }


# ----------------------------------------------------------------------
# 4. Catálogo e Verificação de Prateleira (Autenticado)
# ----------------------------------------------------------------------
@router.get("/{inventory_id}/catalog-cache")
def get_catalog_cache(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    items = (
        db.query(
            InventoryItem.isbn,
            InventoryItem.title,
            InventoryItem.publisher,
            InventoryItem.category,
            InventoryItem.default_location
        )
        .filter(InventoryItem.inventory_id == inventory_id)
        .all()
    )
    return [
        {
            "isbn": it[0],
            "title": it[1],
            "publisher": it[2] or "",
            "category": it[3] or "",
            "default_location": it[4] or "",
        }
        for it in items
    ]


@router.get("/{inventory_id}/sessions/check-location", response_model=inv_schemas.LocationCheckResponse)
def check_location(
    company_id: int,
    inventory_id: int,
    location: str = Query(..., min_length=1),
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    clean_loc = location.strip().upper()
    existing_session = (
        db.query(InventorySession)
        .filter(
            InventorySession.inventory_id == inventory_id,
            func.upper(InventorySession.location) == clean_loc,
            InventorySession.status != SessionStatus.CANCELADA.value
        )
        .order_by(desc(InventorySession.id))
        .first()
    )
    if not existing_session:
        return inv_schemas.LocationCheckResponse(
            location=clean_loc,
            exists=False,
            status=None,
            last_operator_name=None,
            last_operator_id=None,
            last_counted_at=None,
            total_scans_previous=0,
            session_id=None,
        )
    
    op_name = existing_session.operator_name
    if not op_name and existing_session.user_id:
        operator = db.query(user_models.User).filter(user_models.User.id == existing_session.user_id).first()
        op_name = operator.name if operator else None
    if not op_name:
        op_name = f"Operador #{existing_session.id}"
    
    return inv_schemas.LocationCheckResponse(
        location=clean_loc,
        exists=True,
        status=existing_session.status,
        last_operator_name=op_name,
        last_operator_id=existing_session.user_id,
        last_counted_at=existing_session.closed_at or existing_session.started_at,
        total_scans_previous=existing_session.total_scans,
        session_id=existing_session.id,
    )


@router.post("/{inventory_id}/sessions", response_model=inv_schemas.InventorySessionResponse)
def open_session(
    company_id: int,
    inventory_id: int,
    payload: inv_schemas.InventorySessionCreate,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    if inv.status != InventoryStatus.EM_ANDAMENTO.value:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Contagem bloqueada. O inventário não está em andamento (Status atual: '{inv.status}')."
        )
    
    clean_loc = payload.location.strip().upper()
    if not clean_loc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Localização é obrigatória.")
    
    existing_sessions = (
        db.query(InventorySession)
        .filter(
            InventorySession.inventory_id == inventory_id,
            func.upper(InventorySession.location) == clean_loc,
            InventorySession.status != SessionStatus.CANCELADA.value
        )
        .all()
    )

    # Se o operador solicitou "continuar" contagem na prateleira existente
    if payload.mode == "continue" and existing_sessions:
        last_sess = max(existing_sessions, key=lambda s: s.id)
        last_sess.status = SessionStatus.ABERTA.value
        last_sess.closed_at = None
        if payload.operator_name:
            last_sess.operator_name = payload.operator_name
        db.commit()
        db.refresh(last_sess)
        return inv_schemas.InventorySessionResponse(
            id=last_sess.id,
            inventory_id=last_sess.inventory_id,
            company_id=last_sess.company_id,
            user_id=last_sess.user_id,
            user_name=current_user.name,
            operator_name=last_sess.operator_name,
            location=last_sess.location,
            session_type=last_sess.session_type,
            round_number=last_sess.round_number,
            status=last_sess.status,
            total_scans=last_sess.total_scans,
            started_at=last_sess.started_at,
            closed_at=last_sess.closed_at,
        )

    session_type = SessionType.CONTAGEM.value
    round_number = 1
    if existing_sessions:
        session_type = SessionType.RECONTAGEM_AUDITORIA.value
        max_round = max(s.round_number for s in existing_sessions)
        round_number = max_round + 1
    
    new_session = InventorySession(
        inventory_id=inventory_id,
        company_id=company_id,
        user_id=current_user.id,
        operator_name=payload.operator_name or current_user.name,
        location=clean_loc,
        session_type=session_type,
        round_number=round_number,
        status=SessionStatus.ABERTA.value,
        total_scans=0,
    )
    db.add(new_session)
    db.commit()
    db.refresh(new_session)
    
    return inv_schemas.InventorySessionResponse(
        id=new_session.id,
        inventory_id=new_session.inventory_id,
        company_id=new_session.company_id,
        user_id=new_session.user_id,
        user_name=current_user.name,
        operator_name=new_session.operator_name,
        location=new_session.location,
        session_type=new_session.session_type,
        round_number=new_session.round_number,
        status=new_session.status,
        total_scans=new_session.total_scans,
        started_at=new_session.started_at,
        closed_at=new_session.closed_at,
    )


@router.post("/{inventory_id}/sessions/{session_id}/scans/batch", response_model=inv_schemas.InventoryScanBatchResponse)
def sync_scans_batch(
    company_id: int,
    inventory_id: int,
    session_id: int,
    payload: inv_schemas.InventoryScanBatchRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    if inv.status != InventoryStatus.EM_ANDAMENTO.value:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Gravação de bips bloqueada. O inventário não está em andamento (Status atual: '{inv.status}')."
        )
    
    sess = db.query(InventorySession).filter(InventorySession.id == session_id, InventorySession.inventory_id == inventory_id).first()
    if not sess:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sessão não encontrada.")
    if sess.status == SessionStatus.CONCLUIDA.value:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Esta sessão já foi concluída.")
    
    if not payload.scans:
        return inv_schemas.InventoryScanBatchResponse(synced_count=0, ignored_duplicate_count=0, session_total_scans=sess.total_scans)
    
    client_uuids = [s.client_uuid for s in payload.scans]
    existing_uuids = set(
        x[0] for x in db.query(InventoryScan.client_uuid)
        .filter(InventoryScan.session_id == session_id, InventoryScan.client_uuid.in_(client_uuids))
        .all()
    )
    
    synced_count = 0
    ignored_count = 0
    
    for item in payload.scans:
        if item.client_uuid in existing_uuids:
            ignored_count += 1
            continue
        isbn_clean = item.isbn.strip()
        if not isbn_clean:
            continue
        
        item_expected = db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id, InventoryItem.isbn == isbn_clean).first()
        clean_title = (item.title.strip() if item.title and item.title.strip() else f"Item Avulso ({isbn_clean})")
        clean_publisher = (item.publisher.strip() if item.publisher and item.publisher.strip() else None)

        if not item_expected:
            db.add(InventoryItem(
                inventory_id=inventory_id,
                company_id=company_id,
                isbn=isbn_clean,
                title=clean_title,
                publisher=clean_publisher,
                default_location=sess.location,
                is_unregistered=True
            ))
            db.flush()
        elif item_expected.is_unregistered:
            if item.title and item.title.strip() and (item_expected.title.startswith("Item Avulso") or not item_expected.title):
                item_expected.title = item.title.strip()
            if item.publisher and item.publisher.strip() and not item_expected.publisher:
                item_expected.publisher = item.publisher.strip()
        
        new_scan = InventoryScan(
            session_id=session_id,
            inventory_id=inventory_id,
            company_id=company_id,
            user_id=current_user.id,
            operator_name=item.operator_name or current_user.name,
            isbn=isbn_clean,
            location=sess.location,
            quantity=item.quantity if item.quantity > 0 else 1,
            client_uuid=item.client_uuid,
            scanned_at=item.scanned_at,
        )
        db.add(new_scan)
        existing_uuids.add(item.client_uuid)
        synced_count += 1
    
    db.commit()
    new_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
    sess.total_scans = int(new_total or 0)
    db.commit()
    
    return inv_schemas.InventoryScanBatchResponse(
        synced_count=synced_count,
        ignored_duplicate_count=ignored_count,
        session_total_scans=sess.total_scans
    )


@router.put("/{inventory_id}/sessions/{session_id}/close", response_model=inv_schemas.InventorySessionResponse)
def close_session(
    company_id: int,
    inventory_id: int,
    session_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    sess = db.query(InventorySession).filter(InventorySession.id == session_id, InventorySession.inventory_id == inventory_id).first()
    if not sess:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sessão não encontrada.")
    
    sess.status = SessionStatus.CONCLUIDA.value
    sess.closed_at = func.now()
    total_scans = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
    sess.total_scans = int(total_scans or 0)
    db.commit()
    db.refresh(sess)
    
    return inv_schemas.InventorySessionResponse(
        id=sess.id,
        inventory_id=sess.inventory_id,
        company_id=sess.company_id,
        user_id=sess.user_id,
        user_name=current_user.name,
        operator_name=sess.operator_name,
        location=sess.location,
        session_type=sess.session_type,
        round_number=sess.round_number,
        status=sess.status,
        total_scans=sess.total_scans,
        started_at=sess.started_at,
        closed_at=sess.closed_at,
    )


@router.put("/{inventory_id}/finalize", response_model=inv_schemas.InventoryResponse)
def finalize_inventory(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    if inv.status == InventoryStatus.FINALIZADO.value:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Este inventário já se encontra finalizado.")
    
    open_sessions = (
        db.query(InventorySession)
        .filter(InventorySession.inventory_id == inventory_id, InventorySession.status == SessionStatus.ABERTA.value)
        .all()
    )
    for s in open_sessions:
        s.status = SessionStatus.CONCLUIDA.value
        s.closed_at = func.now()
    
    inv.status = InventoryStatus.FINALIZADO.value
    inv.finalized_by_user_id = current_user.id
    inv.finalized_at = func.now()
    db.commit()
    db.refresh(inv)
    
    total_scanned = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.inventory_id == inv.id).scalar()
    total_sessions = db.query(func.count(InventorySession.id)).filter(InventorySession.inventory_id == inv.id).scalar()
    
    return inv_schemas.InventoryResponse(
        id=inv.id,
        company_id=inv.company_id,
        code=inv.code,
        name=inv.name,
        status=inv.status,
        description=inv.description,
        total_expected_skus=inv.total_expected_skus,
        total_scanned_items=int(total_scanned or 0),
        total_sessions=int(total_sessions or 0),
        open_sessions=0,
        access_token=inv.access_token,
        is_public_access_enabled=inv.is_public_access_enabled,
        created_by_user_id=inv.created_by_user_id,
        finalized_by_user_id=inv.finalized_by_user_id,
        finalized_at=inv.finalized_at,
        created_at=inv.created_at,
        updated_at=inv.updated_at,
    )


@router.put("/{inventory_id}/status")
def update_inventory_status(
    company_id: int,
    inventory_id: int,
    payload: inv_schemas.InventoryStatusUpdateRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    # 1. Verificação OBRIGATÓRIA da senha do usuário seller logado
    if not payload.password or not security.verify_password(payload.password, current_user.password_hash):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Senha de usuário incorreta. Operação de alteração de status cancelada."
        )
    
    new_status = payload.status.strip().upper()
    valid_statuses = [s.value for s in InventoryStatus]
    if new_status not in valid_statuses:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Status '{payload.status}' é inválido. Escolha entre: {', '.join(valid_statuses)}."
        )
    
    inv.status = new_status
    
    # 2. Se for alterado para FINALIZADO ou CANCELADO, encerra automaticamente sessões em aberto
    if new_status in [InventoryStatus.FINALIZADO.value, InventoryStatus.CANCELADO.value]:
        open_sessions = (
            db.query(InventorySession)
            .filter(InventorySession.inventory_id == inventory_id, InventorySession.status == SessionStatus.ABERTA.value)
            .all()
        )
        for s in open_sessions:
            s.status = SessionStatus.CONCLUIDA.value if new_status == InventoryStatus.FINALIZADO.value else SessionStatus.CANCELADA.value
            s.closed_at = func.now()
            
        if new_status == InventoryStatus.FINALIZADO.value:
            inv.finalized_by_user_id = current_user.id
            inv.finalized_at = func.now()
            
    db.commit()
    db.refresh(inv)
    
    total_scanned = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.inventory_id == inv.id).scalar()
    total_sessions = db.query(func.count(InventorySession.id)).filter(InventorySession.inventory_id == inv.id).scalar()
    open_count = db.query(func.count(InventorySession.id)).filter(InventorySession.inventory_id == inv.id, InventorySession.status == SessionStatus.ABERTA.value).scalar()
    
    return inv_schemas.InventoryResponse(
        id=inv.id,
        company_id=inv.company_id,
        code=inv.code,
        name=inv.name,
        status=inv.status,
        description=inv.description,
        total_expected_skus=inv.total_expected_skus,
        total_scanned_items=int(total_scanned or 0),
        total_sessions=int(total_sessions or 0),
        open_sessions=int(open_count or 0),
        access_token=inv.access_token,
        is_public_access_enabled=inv.is_public_access_enabled,
        supervisor_pin=inv.supervisor_pin,
        created_by_user_id=inv.created_by_user_id,
        finalized_by_user_id=inv.finalized_by_user_id,
        finalized_at=inv.finalized_at,
        created_at=inv.created_at,
        updated_at=inv.updated_at,
    )


@router.get("/{inventory_id}/sessions-list")
def list_inventory_sessions(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    sessions = (
        db.query(
            InventorySession.id,
            InventorySession.location,
            InventorySession.session_type,
            InventorySession.round_number,
            InventorySession.status,
            InventorySession.total_scans,
            InventorySession.started_at,
            InventorySession.closed_at,
            InventorySession.operator_name,
            user_models.User.name.label("user_name")
        )
        .outerjoin(user_models.User, InventorySession.user_id == user_models.User.id)
        .filter(InventorySession.inventory_id == inventory_id)
        .order_by(InventorySession.location, InventorySession.round_number)
        .all()
    )
    
    return [
        {
            "id": s.id,
            "location": s.location,
            "session_type": s.session_type,
            "round_number": s.round_number,
            "status": s.status,
            "total_scans": s.total_scans,
            "started_at": s.started_at,
            "closed_at": s.closed_at,
            "operator_name": s.operator_name or s.user_name or f"Operador #{s.id}",
        }
        for s in sessions
    ]


@router.get("/{inventory_id}/discrepancies", response_model=List[inv_schemas.DiscrepancyItemResponse])
def get_discrepancies(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    query = text("""
        SELECT 
            s.location,
            s.round_number,
            sc.isbn,
            SUM(sc.quantity) as qty
        FROM inv_inventory_scan sc
        JOIN inv_inventory_session s ON sc.session_id = s.id
        WHERE sc.inventory_id = :inv_id AND s.status != 'CANCELADA'
        GROUP BY s.location, s.round_number, sc.isbn
    """)
    rows = db.execute(query, {"inv_id": inventory_id}).fetchall()
    
    counts_map = {}
    for r in rows:
        loc, round_num, isbn, qty = r[0], r[1], r[2], int(r[3])
        key = (loc, isbn)
        if key not in counts_map:
            counts_map[key] = {1: 0, 2: 0}
        counts_map[key][round_num] = counts_map[key].get(round_num, 0) + qty
        
    items_meta = {
        it.isbn: it for it in db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id).all()
    }
    
    results = []
    for (loc, isbn), rounds in counts_map.items():
        c1 = rounds.get(1, 0)
        c2 = rounds.get(2, 0)
        meta = items_meta.get(isbn)
        diff = c1 - c2
        has_div = (c2 > 0 and c1 != c2)
        val_qty = c2 if c2 > 0 else c1
        
        results.append(inv_schemas.DiscrepancyItemResponse(
            isbn=isbn,
            title=meta.title if meta else f"Item {isbn}",
            publisher=meta.publisher if meta else None,
            category=meta.category if meta else None,
            default_location=meta.default_location if meta else None,
            location=loc,
            count_1_qty=c1,
            count_2_qty=c2,
            difference=diff,
            has_divergence=has_div,
            validated_qty=val_qty
        ))
        
    return sorted(results, key=lambda x: (x.location, x.isbn))


@router.get("/{inventory_id}/sku-summary", response_model=List[inv_schemas.SkuSummaryResponse])
def get_sku_summary(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    
    query_loc_scans = text("""
        SELECT 
            s.location,
            sc.isbn,
            SUM(CASE WHEN s.round_number = 1 THEN sc.quantity ELSE 0 END) as c1,
            SUM(CASE WHEN s.round_number > 1 THEN sc.quantity ELSE 0 END) as c2
        FROM inv_inventory_scan sc
        JOIN inv_inventory_session s ON sc.session_id = s.id
        WHERE sc.inventory_id = :inv_id AND s.status != 'CANCELADA'
        GROUP BY s.location, sc.isbn
    """)
    loc_rows = db.execute(query_loc_scans, {"inv_id": inventory_id}).fetchall()
    
    isbn_locs = {}
    isbn_c1 = {}
    isbn_c2 = {}
    isbn_validated = {}
    
    for r in loc_rows:
        loc, isbn, c1, c2 = r[0], r[1], int(r[2]), int(r[3])
        val_loc = c2 if c2 > 0 else c1
        
        if isbn not in isbn_locs:
            isbn_locs[isbn] = []
            isbn_c1[isbn] = 0
            isbn_c2[isbn] = 0
            isbn_validated[isbn] = 0
            
        isbn_locs[isbn].append(loc)
        isbn_c1[isbn] += c1
        isbn_c2[isbn] += c2
        isbn_validated[isbn] += val_loc

    items = db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id).all()
    items_map = {it.isbn: it for it in items}
    
    all_isbns = set(items_map.keys()).union(set(isbn_validated.keys()))
    
    results = []
    for isbn in all_isbns:
        it = items_map.get(isbn)
        c1 = isbn_c1.get(isbn, 0)
        c2 = isbn_c2.get(isbn, 0)
        val_qty = isbn_validated.get(isbn, 0)
        locs = sorted(list(set(isbn_locs.get(isbn, []))))
        has_div = (c2 > 0 and c1 != c2)
        
        results.append(inv_schemas.SkuSummaryResponse(
            isbn=isbn,
            title=it.title if it else f"Item {isbn}",
            publisher=it.publisher if it else None,
            category=it.category if it else None,
            default_location=it.default_location if it else None,
            locations_list=locs,
            total_count_1=c1,
            total_count_2=c2,
            total_validated_qty=val_qty,
            has_divergence=has_div
        ))
        
    return sorted(results, key=lambda x: x.title)


@router.post("/{inventory_id}/audit-adjust")
def adjust_audit_quantity(
    company_id: int,
    inventory_id: int,
    payload: inv_schemas.AuditAdjustmentRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    expected_pin = inv.supervisor_pin or "1234"
    if payload.pin.strip() != expected_pin:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Senha do Supervisor incorreta.")
    
    loc_clean = payload.location.strip().upper()
    
    sess_query = db.query(InventorySession).filter(
        InventorySession.inventory_id == inventory_id,
        InventorySession.location == loc_clean
    )
    if payload.round_number:
        sess = sess_query.filter(InventorySession.round_number == payload.round_number).first()
    else:
        sess = sess_query.order_by(InventorySession.round_number.desc()).first()
        
    if not sess:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Nenhuma sessão encontrada para a prateleira '{loc_clean}'.")
    
    db.query(InventoryScan).filter(
        InventoryScan.session_id == sess.id,
        InventoryScan.isbn == payload.isbn
    ).delete(synchronize_session=False)
    
    if payload.new_quantity > 0:
        client_uuid = str(uuid.uuid4())
        scan = InventoryScan(
            client_uuid=client_uuid,
            inventory_id=inventory_id,
            session_id=sess.id,
            company_id=company_id,
            user_id=current_user.id,
            isbn=payload.isbn,
            location=loc_clean,
            quantity=payload.new_quantity,
            scanned_at=datetime.utcnow()
        )
        db.add(scan)
        
    db.commit()
    
    new_sess_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == sess.id).scalar()
    sess.total_scans = int(new_sess_total)
    
    new_inv_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.inventory_id == inventory_id).scalar()
    inv.total_scanned_items = int(new_inv_total)
    db.commit()
    
    return {
        "message": f"Quantidade do ISBN {payload.isbn} na prateleira {loc_clean} ajustada para {payload.new_quantity} un!",
        "new_session_total": sess.total_scans,
        "new_inventory_total": inv.total_scanned_items
    }


@router.put("/{inventory_id}/items/{isbn}")
def update_inventory_item_meta(
    company_id: int,
    inventory_id: int,
    isbn: str,
    payload: inv_schemas.InventoryItemUpdateRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    item = db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id, InventoryItem.isbn == isbn).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Item não encontrado no inventário.")
    
    if payload.title is not None and payload.title.strip():
        item.title = payload.title.strip()
    if payload.publisher is not None:
        item.publisher = payload.publisher.strip()
    if payload.category is not None:
        item.category = payload.category.strip()
        
    db.commit()
    return {"message": f"Dados do produto {isbn} atualizados com sucesso!", "title": item.title, "publisher": item.publisher}


@router.get("/{inventory_id}/export-excel")
def export_inventory_excel(
    company_id: int,
    inventory_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Inventário não encontrado.")
    
    wb = openpyxl.Workbook()
    
    # Aba 1 (PRINCIPAL): Saldos Validados (Título, ISBN, Marca, Quantidade)
    ws_val = wb.active
    ws_val.title = "Saldos Validados"
    ws_val.append([
        "Título",
        "ISBN",
        "Marca",
        "Quantidade Validada"
    ])
    
    # Query de contagens por localização e rodada
    query_loc_scans = text("""
        SELECT 
            s.location,
            sc.isbn,
            SUM(CASE WHEN s.round_number = 1 THEN sc.quantity ELSE 0 END) as c1,
            SUM(CASE WHEN s.round_number > 1 THEN sc.quantity ELSE 0 END) as c2
        FROM inv_inventory_scan sc
        JOIN inv_inventory_session s ON sc.session_id = s.id
        WHERE sc.inventory_id = :inv_id AND s.status != 'CANCELADA'
        GROUP BY s.location, sc.isbn
    """)
    loc_rows = db.execute(query_loc_scans, {"inv_id": inventory_id}).fetchall()
    
    # Calcula saldo validado por ISBN (Recontagem se houver, senão 1ª Contagem)
    isbn_validated_totals = {}
    for r in loc_rows:
        isbn, c1, c2 = r[1], int(r[2]), int(r[3])
        val_loc = c2 if c2 > 0 else c1
        isbn_validated_totals[isbn] = isbn_validated_totals.get(isbn, 0) + val_loc

    items = db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id).order_by(InventoryItem.title).all()
    
    for it in items:
        val_qty = isbn_validated_totals.get(it.isbn, 0)
        marca = it.publisher or it.category or "N/A"
        ws_val.append([
            it.title,
            it.isbn,
            marca,
            val_qty
        ])
        
    # Aba 2: Consolidado Geral
    ws1 = wb.create_sheet(title="Consolidado Geral")
    ws1.append([
        "ISBN / Código de Barras",
        "Título",
        "Editora / Marca",
        "Categoria",
        "Endereço Padrão",
        "Qtd 1ª Contagem",
        "Qtd Recontagem/Auditoria",
        "Divergência",
        "Total Geral Apurado",
        "Saldo Validado Final",
        "Status Cadastral"
    ])
    
    query_isbn = text("""
        SELECT 
            sc.isbn,
            SUM(CASE WHEN s.round_number = 1 THEN sc.quantity ELSE 0 END) as c1,
            SUM(CASE WHEN s.round_number > 1 THEN sc.quantity ELSE 0 END) as c2,
            SUM(sc.quantity) as total_scanned
        FROM inv_inventory_scan sc
        JOIN inv_inventory_session s ON sc.session_id = s.id
        WHERE sc.inventory_id = :inv_id AND s.status != 'CANCELADA'
        GROUP BY sc.isbn
    """)
    scan_rows = {r[0]: (int(r[1]), int(r[2]), int(r[3])) for r in db.execute(query_isbn, {"inv_id": inventory_id}).fetchall()}
    
    for it in items:
        c1, c2, total = scan_rows.get(it.isbn, (0, 0, 0))
        diff = c1 - c2 if c2 > 0 else 0
        val_qty = isbn_validated_totals.get(it.isbn, 0)
        status_cad = "Item Novo (Fora da Base)" if it.is_unregistered else "Cadastrado"
        ws1.append([
            it.isbn, it.title, it.publisher or "", it.category or "", it.default_location or "",
            c1, c2, diff, total, val_qty, status_cad
        ])
        
    # Aba 3: Detalhamento por Prateleira
    ws2 = wb.create_sheet(title="Detalhamento por Prateleira")
    ws2.append([
        "ID Sessão", "Prateleira / Localização", "Tipo de Contagem", "Rodada", "Status",
        "Operador", "Início", "Término", "Total Peças Bipadas"
    ])
    
    sessions = (
        db.query(
            InventorySession.id,
            InventorySession.location,
            InventorySession.session_type,
            InventorySession.round_number,
            InventorySession.status,
            InventorySession.operator_name,
            user_models.User.name.label("user_name"),
            InventorySession.started_at,
            InventorySession.closed_at,
            InventorySession.total_scans
        )
        .outerjoin(user_models.User, InventorySession.user_id == user_models.User.id)
        .filter(InventorySession.inventory_id == inventory_id)
        .order_by(InventorySession.location, InventorySession.round_number)
        .all()
    )
    
    for s in sessions:
        op = s.operator_name or s.user_name or f"Operador #{s.id}"
        ws2.append([
            s.id, s.location, s.session_type, s.round_number, s.status, op,
            s.started_at.strftime("%d/%m/%Y %H:%M:%S") if s.started_at else "",
            s.closed_at.strftime("%d/%m/%Y %H:%M:%S") if s.closed_at else "",
            s.total_scans
        ])
        
    # Aba 4: Itens Fora da Base
    ws3 = wb.create_sheet(title="Itens Fora da Base")
    ws3.append(["ISBN / Código", "Título", "Marca / Editora", "Localização Bipada", "Qtd Validada"])
    
    unreg_items = db.query(InventoryItem).filter(InventoryItem.inventory_id == inventory_id, InventoryItem.is_unregistered == True).all()
    for un in unreg_items:
        val_qty = isbn_validated_totals.get(un.isbn, 0)
        ws3.append([un.isbn, un.title, un.publisher or "", un.default_location or "", val_qty])
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"Relatorio_Inventario_{inv.code}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ======================================================================
# ROTAS PÚBLICAS / LINK DIRETO DO OPERADOR (SEM LOGIN)
# ======================================================================

@public_router.get("/{access_token}", response_model=inv_schemas.PublicInventoryInfo)
def get_public_inventory_info(access_token: str, db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    company = db.query(Company).filter(Company.id == inv.company_id).first()
    return inv_schemas.PublicInventoryInfo(
        id=inv.id,
        company_id=inv.company_id,
        company_name=company.name if company else "Empresa",
        code=inv.code,
        name=inv.name,
        description=inv.description,
        status=inv.status,
        total_expected_skus=inv.total_expected_skus,
    )


@public_router.get("/{access_token}/catalog-cache")
def get_public_catalog_cache(access_token: str, db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    items = (
        db.query(
            InventoryItem.isbn,
            InventoryItem.title,
            InventoryItem.publisher,
            InventoryItem.category,
            InventoryItem.default_location
        )
        .filter(InventoryItem.inventory_id == inv.id)
        .all()
    )
    return [
        {
            "isbn": it[0],
            "title": it[1],
            "publisher": it[2] or "",
            "category": it[3] or "",
            "default_location": it[4] or "",
        }
        for it in items
    ]


@public_router.get("/{access_token}/check-location", response_model=inv_schemas.LocationCheckResponse)
def check_public_location(access_token: str, location: str = Query(..., min_length=1), db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    clean_loc = location.strip().upper()
    existing_session = (
        db.query(InventorySession)
        .filter(
            InventorySession.inventory_id == inv.id,
            func.upper(InventorySession.location) == clean_loc,
            InventorySession.status != SessionStatus.CANCELADA.value
        )
        .order_by(desc(InventorySession.id))
        .first()
    )
    if not existing_session:
        return inv_schemas.LocationCheckResponse(
            location=clean_loc,
            exists=False,
            status=None,
            last_operator_name=None,
            last_operator_id=None,
            last_counted_at=None,
            total_scans_previous=0,
            session_id=None,
        )
    
    op_name = existing_session.operator_name
    if not op_name and existing_session.user_id:
        user = db.query(user_models.User).filter(user_models.User.id == existing_session.user_id).first()
        if user:
            op_name = user.name
    if not op_name:
        op_name = f"Operador #{existing_session.id}"
        
    return inv_schemas.LocationCheckResponse(
        location=clean_loc,
        exists=True,
        status=existing_session.status,
        last_operator_name=op_name,
        last_operator_id=existing_session.user_id,
        last_counted_at=existing_session.closed_at or existing_session.started_at,
        total_scans_previous=existing_session.total_scans,
        session_id=existing_session.id,
    )


@public_router.post("/{access_token}/sessions", response_model=inv_schemas.InventorySessionResponse)
def open_public_session(access_token: str, payload: inv_schemas.PublicSessionCreate, db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    if inv.status != InventoryStatus.EM_ANDAMENTO.value:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Contagem bloqueada. O inventário não está em andamento (Status atual: '{inv.status}')."
        )
    clean_loc = payload.location.strip().upper()
    clean_op = payload.operator_name.strip()
    if not clean_loc:
        raise HTTPException(status_code=400, detail="Localização é obrigatória.")
    if not clean_op:
        raise HTTPException(status_code=400, detail="Nome do operador é obrigatório.")

    existing_sessions = (
        db.query(InventorySession)
        .filter(
            InventorySession.inventory_id == inv.id,
            func.upper(InventorySession.location) == clean_loc,
            InventorySession.status != SessionStatus.CANCELADA.value
        )
        .all()
    )

    if payload.mode == "continue" and existing_sessions:
        last_sess = max(existing_sessions, key=lambda s: s.id)
        last_sess.status = SessionStatus.ABERTA.value
        last_sess.closed_at = None
        if clean_op:
            last_sess.operator_name = clean_op
        db.commit()
        db.refresh(last_sess)
        return inv_schemas.InventorySessionResponse(
            id=last_sess.id,
            inventory_id=last_sess.inventory_id,
            company_id=last_sess.company_id,
            user_id=None,
            user_name=clean_op,
            operator_name=last_sess.operator_name,
            location=last_sess.location,
            session_type=last_sess.session_type,
            round_number=last_sess.round_number,
            status=last_sess.status,
            total_scans=last_sess.total_scans,
            started_at=last_sess.started_at,
            closed_at=last_sess.closed_at,
        )

    session_type = SessionType.CONTAGEM.value
    round_number = 1
    if existing_sessions:
        session_type = SessionType.RECONTAGEM_AUDITORIA.value
        max_round = max(s.round_number for s in existing_sessions)
        round_number = max_round + 1

    new_session = InventorySession(
        inventory_id=inv.id,
        company_id=inv.company_id,
        user_id=None,
        operator_name=clean_op,
        location=clean_loc,
        session_type=session_type,
        round_number=round_number,
        status=SessionStatus.ABERTA.value,
        total_scans=0,
    )
    db.add(new_session)
    db.commit()
    db.refresh(new_session)

    return inv_schemas.InventorySessionResponse(
        id=new_session.id,
        inventory_id=new_session.inventory_id,
        company_id=new_session.company_id,
        user_id=None,
        user_name=clean_op,
        operator_name=clean_op,
        location=new_session.location,
        session_type=new_session.session_type,
        round_number=new_session.round_number,
        status=new_session.status,
        total_scans=new_session.total_scans,
        started_at=new_session.started_at,
        closed_at=new_session.closed_at,
    )


@public_router.post("/{access_token}/sessions/{session_id}/scans/batch", response_model=inv_schemas.InventoryScanBatchResponse)
def sync_public_scans_batch(access_token: str, session_id: int, payload: inv_schemas.InventoryScanBatchRequest, db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    if inv.status != InventoryStatus.EM_ANDAMENTO.value:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Gravação de bips bloqueada. O inventário não está em andamento (Status atual: '{inv.status}')."
        )
    sess = db.query(InventorySession).filter(InventorySession.id == session_id, InventorySession.inventory_id == inv.id).first()
    if not sess:
        raise HTTPException(status_code=404, detail="Sessão não encontrada.")
    if sess.status == SessionStatus.CONCLUIDA.value:
        raise HTTPException(status_code=409, detail="Esta sessão já foi concluída e não aceita mais leituras.")
    
    if not payload.scans:
        return inv_schemas.InventoryScanBatchResponse(synced_count=0, ignored_duplicate_count=0, session_total_scans=sess.total_scans)

    client_uuids = [s.client_uuid for s in payload.scans]
    existing_uuids = set(
        x[0] for x in db.query(InventoryScan.client_uuid)
        .filter(InventoryScan.session_id == session_id, InventoryScan.client_uuid.in_(client_uuids))
        .all()
    )

    synced_count = 0
    ignored_count = 0
    for item in payload.scans:
        if item.client_uuid in existing_uuids:
            ignored_count += 1
            continue
        isbn_clean = item.isbn.strip()
        if not isbn_clean:
            continue
        
        item_expected = db.query(InventoryItem).filter(InventoryItem.inventory_id == inv.id, InventoryItem.isbn == isbn_clean).first()
        clean_title = (item.title.strip() if item.title and item.title.strip() else f"Item Avulso ({isbn_clean})")
        clean_publisher = (item.publisher.strip() if item.publisher and item.publisher.strip() else None)

        if not item_expected:
            db.add(InventoryItem(
                inventory_id=inv.id,
                company_id=inv.company_id,
                isbn=isbn_clean,
                title=clean_title,
                publisher=clean_publisher,
                default_location=sess.location,
                is_unregistered=True
            ))
            db.flush()
        elif item_expected.is_unregistered:
            if item.title and item.title.strip() and (item_expected.title.startswith("Item Avulso") or not item_expected.title):
                item_expected.title = item.title.strip()
            if item.publisher and item.publisher.strip() and not item_expected.publisher:
                item_expected.publisher = item.publisher.strip()

        db.add(InventoryScan(
            session_id=session_id,
            inventory_id=inv.id,
            company_id=inv.company_id,
            user_id=None,
            operator_name=item.operator_name or sess.operator_name,
            isbn=isbn_clean,
            location=sess.location,
            quantity=item.quantity if item.quantity > 0 else 1,
            client_uuid=item.client_uuid,
            scanned_at=item.scanned_at,
        ))
        existing_uuids.add(item.client_uuid)
        synced_count += 1

    db.commit()
    new_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
    sess.total_scans = int(new_total or 0)
    db.commit()

    return inv_schemas.InventoryScanBatchResponse(
        synced_count=synced_count,
        ignored_duplicate_count=ignored_count,
        session_total_scans=sess.total_scans
    )


@public_router.put("/{access_token}/sessions/{session_id}/close", response_model=inv_schemas.InventorySessionResponse)
def close_public_session(access_token: str, session_id: int, db: Session = Depends(get_db)):
    inv = _get_inventory_by_token(access_token, db)
    sess = db.query(InventorySession).filter(InventorySession.id == session_id, InventorySession.inventory_id == inv.id).first()
    if not sess:
        raise HTTPException(status_code=404, detail="Sessão não encontrada.")
    sess.status = SessionStatus.CONCLUIDA.value
    sess.closed_at = func.now()
    total_scans = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
    sess.total_scans = int(total_scans or 0)
    db.commit()
    db.refresh(sess)

    return inv_schemas.InventorySessionResponse(
        id=sess.id,
        inventory_id=sess.inventory_id,
        company_id=sess.company_id,
        user_id=None,
        user_name=sess.operator_name,
        operator_name=sess.operator_name,
        location=sess.location,
        session_type=sess.session_type,
        round_number=sess.round_number,
        status=sess.status,
        total_scans=sess.total_scans,
        started_at=sess.started_at,
        closed_at=sess.closed_at,
    )


# ----------------------------------------------------------------------
# 13. Endpoints de Manutenção & Senha de Supervisor (PIN)
# ----------------------------------------------------------------------

# Validação de PIN (Autenticado)
@router.post("/{inventory_id}/verify-pin", response_model=inv_schemas.PinVerifyResponse)
def verify_inventory_pin(
    company_id: int,
    inventory_id: int,
    payload: inv_schemas.PinVerifyRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Inventário não encontrado.")
    
    expected_pin = inv.supervisor_pin or "1234"
    if payload.pin.strip() != expected_pin:
        raise HTTPException(status_code=401, detail="Senha do Supervisor incorreta.")
    
    return inv_schemas.PinVerifyResponse(success=True, message="Senha do Supervisor autorizada!")


# Validação de PIN (Público via Token)
@public_router.post("/{access_token}/verify-pin", response_model=inv_schemas.PinVerifyResponse)
def verify_public_inventory_pin(
    access_token: str,
    payload: inv_schemas.PinVerifyRequest,
    db: Session = Depends(get_db)
):
    inv = _get_inventory_by_token(access_token, db)
    expected_pin = inv.supervisor_pin or "1234"
    if payload.pin.strip() != expected_pin:
        raise HTTPException(status_code=401, detail="Senha do Supervisor incorreta.")
    
    return inv_schemas.PinVerifyResponse(success=True, message="Senha do Supervisor autorizada!")


# Listar Itens Agregados por Sessão (Autenticado)
@router.get("/{inventory_id}/sessions/{session_id}/items", response_model=List[inv_schemas.SessionItemSummary])
def get_session_items_summary(
    company_id: int,
    inventory_id: int,
    session_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    return _fetch_session_items_summary_db(inventory_id, session_id, db)


# Listar Itens Agregados por Sessão (Público)
@public_router.get("/{access_token}/sessions/{session_id}/items", response_model=List[inv_schemas.SessionItemSummary])
def get_public_session_items_summary(
    access_token: str,
    session_id: int,
    db: Session = Depends(get_db),
):
    inv = _get_inventory_by_token(access_token, db)
    return _fetch_session_items_summary_db(inv.id, session_id, db)


def _fetch_session_items_summary_db(inventory_id: int, session_id: int, db: Session) -> List[inv_schemas.SessionItemSummary]:
    scans_sub = (
        db.query(
            InventoryScan.isbn,
            func.sum(InventoryScan.quantity).label("total_quantity"),
            func.max(InventoryScan.scanned_at).label("last_scanned_at")
        )
        .filter(InventoryScan.inventory_id == inventory_id, InventoryScan.session_id == session_id)
        .group_by(InventoryScan.isbn)
        .all()
    )

    if not scans_sub:
        return []

    isbns = [s.isbn for s in scans_sub]
    catalog_items = (
        db.query(InventoryItem)
        .filter(InventoryItem.inventory_id == inventory_id, InventoryItem.isbn.in_(isbns))
        .all()
    )
    cat_map = {ci.isbn: ci for ci in catalog_items}

    result = []
    for row in scans_sub:
        ci = cat_map.get(row.isbn)
        title = ci.title if ci else f"Item ({row.isbn})"
        publisher = ci.publisher if ci else None
        category = ci.category if ci else None
        
        result.append(inv_schemas.SessionItemSummary(
            isbn=row.isbn,
            title=title,
            publisher=publisher,
            category=category,
            total_quantity=int(row.total_quantity or 0),
            last_scanned_at=row.last_scanned_at
        ))

    result.sort(key=lambda x: x.last_scanned_at or datetime.min, reverse=True)
    return result


# Atualizar Quantidade de um Item na Sessão (com PIN do Supervisor)
@router.put("/{inventory_id}/sessions/{session_id}/items/{isbn}")
def update_session_item_quantity(
    company_id: int,
    inventory_id: int,
    session_id: int,
    isbn: str,
    payload: inv_schemas.SessionItemUpdateRequest,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Inventário não encontrado.")
    return _apply_session_item_update(inv, session_id, isbn, payload.pin, payload.new_quantity, db)


@public_router.put("/{access_token}/sessions/{session_id}/items/{isbn}")
def update_public_session_item_quantity(
    access_token: str,
    session_id: int,
    isbn: str,
    payload: inv_schemas.SessionItemUpdateRequest,
    db: Session = Depends(get_db),
):
    inv = _get_inventory_by_token(access_token, db)
    return _apply_session_item_update(inv, session_id, isbn, payload.pin, payload.new_quantity, db)


def _apply_session_item_update(inv: Inventory, session_id: int, isbn: str, pin: str, new_qty: int, db: Session):
    expected_pin = inv.supervisor_pin or "1234"
    if pin.strip() != expected_pin:
        raise HTTPException(status_code=401, detail="Senha do Supervisor incorreta.")
    
    if new_qty < 0:
        raise HTTPException(status_code=400, detail="Quantidade não pode ser negativa.")

    sess = db.query(InventorySession).filter(InventorySession.id == session_id, InventorySession.inventory_id == inv.id).first()
    if not sess:
        raise HTTPException(status_code=404, detail="Sessão não encontrada.")

    # Remove bips existentes deste ISBN nesta sessão
    db.query(InventoryScan).filter(
        InventoryScan.inventory_id == inv.id,
        InventoryScan.session_id == session_id,
        InventoryScan.isbn == isbn.strip()
    ).delete()

    # Se a nova quantidade for maior que zero, cria um registro consolidado
    if new_qty > 0:
        import uuid
        db.add(InventoryScan(
            inventory_id=inv.id,
            company_id=inv.company_id,
            session_id=session_id,
            operator_name=sess.operator_name or "Supervisor",
            isbn=isbn.strip(),
            location=sess.location,
            quantity=new_qty,
            client_uuid=f"maint-{uuid.uuid4()}",
            scanned_at=func.now()
        ))

    db.commit()

    # Recalcula total da sessão
    new_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
    sess.total_scans = int(new_total or 0)
    db.commit()

    return {"message": "Manutenção realizada com sucesso!", "new_total_scans": sess.total_scans}


# Excluir Item da Sessão (com PIN do Supervisor)
@router.delete("/{inventory_id}/sessions/{session_id}/items/{isbn}")
def delete_session_item(
    company_id: int,
    inventory_id: int,
    session_id: int,
    isbn: str,
    pin: str,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Inventário não encontrado.")
    return _apply_session_item_update(inv, session_id, isbn, pin, 0, db)


@public_router.delete("/{access_token}/sessions/{session_id}/items/{isbn}")
def delete_public_session_item(
    access_token: str,
    session_id: int,
    isbn: str,
    pin: str,
    db: Session = Depends(get_db),
):
    inv = _get_inventory_by_token(access_token, db)
    return _apply_session_item_update(inv, session_id, isbn, pin, 0, db)


# Desfazer Último Bip da Sessão
@router.delete("/{inventory_id}/sessions/{session_id}/scans/last")
def undo_last_scan(
    company_id: int,
    inventory_id: int,
    session_id: int,
    db: Session = Depends(get_db),
    current_user: user_models.User = Depends(dependencies.get_current_user),
):
    _assert_inventory_access(current_user, company_id, db)
    inv = db.query(Inventory).filter(Inventory.id == inventory_id, Inventory.company_id == company_id).first()
    if not inv:
        raise HTTPException(status_code=404, detail="Inventário não encontrado.")
    return _apply_undo_last_scan(inv, session_id, db)


@public_router.delete("/{access_token}/sessions/{session_id}/scans/last")
def undo_public_last_scan(
    access_token: str,
    session_id: int,
    db: Session = Depends(get_db),
):
    inv = _get_inventory_by_token(access_token, db)
    return _apply_undo_last_scan(inv, session_id, db)


def _apply_undo_last_scan(inv: Inventory, session_id: int, db: Session):
    last_scan = (
        db.query(InventoryScan)
        .filter(InventoryScan.inventory_id == inv.id, InventoryScan.session_id == session_id)
        .order_by(InventoryScan.id.desc())
        .first()
    )
    if not last_scan:
        raise HTTPException(status_code=404, detail="Nenhum bip encontrado nesta sessão para desfazer.")

    removed_isbn = last_scan.isbn
    removed_qty = last_scan.quantity
    db.delete(last_scan)
    db.commit()

    sess = db.query(InventorySession).filter(InventorySession.id == session_id).first()
    if sess:
        new_total = db.query(func.coalesce(func.sum(InventoryScan.quantity), 0)).filter(InventoryScan.session_id == session_id).scalar()
        sess.total_scans = int(new_total or 0)
        db.commit()

    return {
        "message": f"Último bip do ISBN {removed_isbn} ({removed_qty} un) foi desfeito!",
        "removed_isbn": removed_isbn,
        "removed_quantity": removed_qty,
        "session_total_scans": sess.total_scans if sess else 0
    }
